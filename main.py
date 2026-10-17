from __future__ import annotations

import asyncio
import csv
import errno
import hashlib
import logging
import os
import re
import sys
import tempfile
from bisect import bisect_left
from collections.abc import Callable
from copy import copy
from itertools import count
from pathlib import Path
from urllib.parse import unquote, urlsplit
from xml.etree.ElementTree import ParseError
from zipfile import BadZipFile

import flet as ft

# This column is used to locate the crop image on disk; it isn't shown as a form field.
IMAGE_LINK_FIELD = "Cropped Box Image Link"
# Flags get their own add/remove chip UI instead of a generic TextField.
FLAGS_FIELD = "Review Flags"
# Reviewed state is a Checkbox, not a TextField.
REVIEWED_FIELD = "Reviewed"
# Columns handled by dedicated UI rather than the generic text-field loop.
CUSTOM_FIELDS = {IMAGE_LINK_FIELD, FLAGS_FIELD, REVIEWED_FIELD}
# These fields tend to hold longer free text, so give them more room to wrap.
MULTILINE_FIELDS = {"Description", "Unit Type"}
# VLM-computed metadata shown for context, not meant to be hand-edited.
READ_ONLY_FIELDS = {"Transcription Confidence (%)"}
# Separator used to pack the flags list into the CSV's single text column.
FLAGS_SEPARATOR = "; "
# CSV stores this as the literal strings "True"/"False", not real booleans.
REVIEWED_TRUE = "True"
REVIEWED_FALSE = "False"

DUPLICATE_FLAG = "Potential Duplicate"
MISSING_BOX_FLAG = "Missing Box"
LOW_CONFIDENCE_FLAG = "Low confidence"
MISSING_IDENTIFIER_FLAG = "Missing identifier"
AMBIGUOUS_IDENTIFIER_FLAG = "Ambiguous identifier"
MULTIPLE_IDENTIFIERS_FLAG = "Multiple identifiers"
MISSING_UNIT_FLAG = "Missing unit"
INVALID_UNIT_FLAG = "Invalid unit"
TRANSCRIPTION_MISMATCH_FLAG = "Transcription mismatch"
FALLBACK_USED_FLAG = "Fallback used"

# The only flags choosable from the "+ Add flag" menu, in menu order.
ALL_FLAGS = [
    DUPLICATE_FLAG,
    MISSING_BOX_FLAG,
    LOW_CONFIDENCE_FLAG,
    MISSING_IDENTIFIER_FLAG,
    AMBIGUOUS_IDENTIFIER_FLAG,
    MULTIPLE_IDENTIFIERS_FLAG,
    MISSING_UNIT_FLAG,
    INVALID_UNIT_FLAG,
    TRANSCRIPTION_MISMATCH_FLAG,
    FALLBACK_USED_FLAG,
]

# (chip background, chip text/delete-icon color) per flag.
FLAG_COLORS: dict[str, tuple[str, str]] = {
    DUPLICATE_FLAG: (ft.Colors.PURPLE, ft.Colors.WHITE),
    MISSING_BOX_FLAG: (ft.Colors.RED_700, ft.Colors.WHITE),
    LOW_CONFIDENCE_FLAG: (ft.Colors.AMBER, ft.Colors.BLACK),
    MISSING_IDENTIFIER_FLAG: (ft.Colors.DEEP_ORANGE_900, ft.Colors.WHITE),
    AMBIGUOUS_IDENTIFIER_FLAG: (ft.Colors.ORANGE, ft.Colors.BLACK),
    MULTIPLE_IDENTIFIERS_FLAG: (ft.Colors.PINK_700, ft.Colors.WHITE),
    MISSING_UNIT_FLAG: (ft.Colors.BROWN, ft.Colors.WHITE),
    INVALID_UNIT_FLAG: (ft.Colors.DEEP_PURPLE, ft.Colors.WHITE),
    TRANSCRIPTION_MISMATCH_FLAG: (ft.Colors.INDIGO, ft.Colors.WHITE),
    FALLBACK_USED_FLAG: (ft.Colors.BLUE_GREY_600, ft.Colors.WHITE),
}
# Covers any flag already in a catalogue but outside the list above (e.g. from
# an older run) — still shown and removable, just not offered when adding.
DEFAULT_FLAG_COLOR = (ft.Colors.GREY_400, ft.Colors.BLACK)

# Catalogue formats the CLI can produce.
CATALOGUE_EXTENSIONS = (".csv", ".xlsx")
# Row-dict key (not a column) remembering which xlsx cells held numbers, so they
# are written back as numbers rather than text.
XLSX_NUMERIC_KEY = "\x00numeric"
# Keep the source worksheet position so discarding a record handles blank rows.
XLSX_ROW_KEY = "\x00row"
# Native Excel hyperlinks are URL-encoded; ordinary paths and CSV formulas aren't.
XLSX_IMAGE_URL_KEY = "\x00image_url"
CatalogueRow = dict[str, object]

# Fixed width of the left-hand form column; the image takes whatever's left.
FORM_WIDTH = 380
COMPACT_WIDTH = 850
PAGE_PADDING = 20
THEME_PREFERENCE = "uma_bct_gui.theme_mode"
logger = logging.getLogger(__name__)

# Only literal arguments are accepted. No spreadsheet formula is evaluated.
HYPERLINK_FORMULA = re.compile(
    r'\s*=\s*HYPERLINK\s*\(\s*"((?:[^"]|"")*)"\s*'
    r'(?:[,;]\s*"(?:[^"]|"")*"\s*)?\)\s*',
    re.IGNORECASE,
)


def action_button(label: str, icon: ft.IconData, bgcolor: ft.ColorValue, color: ft.ColorValue) -> ft.Button:
    """A rounded button whose custom colours switch to Material's muted
    disabled look. Button.bgcolor/color apply in every state, so the colours
    have to be given per ControlState instead."""
    disabled_bg = ft.Colors.with_opacity(0.3, bgcolor)
    disabled_fg = ft.Colors.with_opacity(0.3, color)
    fg = {ft.ControlState.DEFAULT: color, ft.ControlState.DISABLED: disabled_fg}
    return ft.Button(
        content=label,
        icon=icon,
        style=ft.ButtonStyle(
            shape=ft.RoundedRectangleBorder(radius=10),
            bgcolor={ft.ControlState.DEFAULT: bgcolor, ft.ControlState.DISABLED: disabled_bg},
            color=fg,
            icon_color=fg,
        ),
    )


def parse_flags(value: str | None) -> list[str]:
    return [p.strip() for p in (value or "").split(";") if p.strip()]


class CatalogueError(Exception):
    """A catalogue file can't be read or written; the message is user-facing."""


async def open_in_default_app(path: Path) -> None:
    """Ask the desktop to open a file, reporting immediate launch failures."""
    filename = str(path.resolve())
    if sys.platform == "win32":
        await asyncio.to_thread(os.startfile, filename)
        return
    opener = "/usr/bin/open" if sys.platform == "darwin" else "xdg-open"
    process = await asyncio.create_subprocess_exec(
        opener, filename,
        stdin=asyncio.subprocess.DEVNULL,
        stdout=asyncio.subprocess.DEVNULL,
        stderr=asyncio.subprocess.DEVNULL,
    )
    try:
        returncode = await asyncio.wait_for(process.wait(), timeout=5)
    except TimeoutError:
        # Some Linux file associations keep the opener running with the app.
        # Leave that app running and allow the reviewer to keep using the GUI.
        return
    if returncode != 0:
        raise OSError(f"The system file opener exited with status {returncode}.")


def resolve_image_path(folder: Path, reference: str, *, url_encoded: bool = False) -> Path | None:
    """Resolve a local image reference without changing its catalogue value."""
    if not reference:
        return None
    target = reference
    if reference.lstrip().startswith("="):
        match = HYPERLINK_FORMULA.fullmatch(reference)
        if match is None:
            raise ValueError("The image link needs a HYPERLINK formula with a literal file path.")
        target = match.group(1).replace('""', '"')
        url_encoded = False
    if not target:
        return None
    target = target.replace("\\", "/")
    windows_drive = bool(re.match(r"^[A-Za-z]:/", target))
    if not windows_drive:
        url = urlsplit(target)
        if url.scheme and url.scheme.lower() != "file":
            # A colon is legal in a POSIX filename. Keep an existing local file
            # working when its bare name could be mistaken for a URI scheme.
            if "://" in target or not (folder / target).is_file():
                raise ValueError("Only local image files are supported.")
        elif url.scheme:
            target = unquote(url.path)
            if url.netloc and url.netloc.lower() != "localhost":
                if os.name != "nt":
                    raise ValueError("This image link refers to a Windows network share.")
                target = f"//{url.netloc}{target}"
            if os.name == "nt" and re.match(r"^/[A-Za-z]:/", target):
                target = target[1:]
            url_encoded = False
    if url_encoded:
        decoded = unquote(target)
        # Some workbooks store native hyperlink targets as literal paths. Prefer
        # the URL interpretation, but retain a literal '%' filename if it is the
        # only existing target. Explicit file: URLs were already decoded above.
        if decoded == target or (folder / decoded).exists() or not (folder / target).is_file():
            target = decoded
    if "\x00" in target:
        raise ValueError("The image path contains a null character.")
    if os.name != "nt" and (windows_drive or target.startswith("//")):
        raise ValueError("This image link refers to a Windows drive or network share.")
    return folder / target


def find_catalogue(folder: Path) -> Path | None:
    """The catalogue.csv / catalogue.xlsx in folder (the newer one if both exist)."""
    try:
        existing = []
        names = {f"catalogue{ext}" for ext in CATALOGUE_EXTENSIONS}
        for path in folder.iterdir():
            if path.name.casefold() in names and path.is_file():
                if any(other.suffix.casefold() == path.suffix.casefold() for other in existing):
                    raise CatalogueError(
                        f"More than one {path.suffix.lower()} catalogue has the same name with different casing. "
                        "Keep one catalogue of each format in the folder."
                    )
                existing.append(path)
        return max(
            existing, key=lambda p: (p.stat().st_mtime_ns, p.suffix.lower() == ".csv")
        ) if existing else None
    except OSError as ex:
        raise CatalogueError(f"Could not read the selected folder: {ex}") from ex


def catalogue_revision(path: Path) -> str:
    """Hash the file contents so a save can detect edits made by another program."""
    try:
        with path.open("rb") as source:
            return hashlib.file_digest(source, "sha256").hexdigest()
    except OSError as ex:
        raise CatalogueError(f"Could not read {path.name}: {ex}") from ex


def _validate_headers(fieldnames: list[str]) -> None:
    if not fieldnames or any(not name.strip() or name.startswith("\x00") for name in fieldnames):
        raise CatalogueError("The catalogue needs a non-empty name for every column.")
    if len(set(fieldnames)) != len(fieldnames):
        raise CatalogueError("The catalogue contains duplicate column names.")


def load_rows(path: Path) -> tuple[list[str], list[CatalogueRow]]:
    """Read a catalogue into (header order, row dicts of strings)."""
    try:
        return _load_xlsx(path) if path.suffix.lower() == ".xlsx" else _load_csv(path)
    except (OSError, UnicodeError, csv.Error, BadZipFile, ParseError, ValueError, KeyError, EOFError) as ex:
        raise CatalogueError(f"Could not read {path.name}: {ex}") from ex


def save_rows(
    path: Path,
    fieldnames: list[str],
    rows: list[CatalogueRow],
    *,
    expected_revision: str | None = None,
) -> str:
    """Save beside the original, then replace it only if its contents have not changed."""
    _validate_headers(fieldnames)
    temporary_path: Path | None = None
    try:
        original_revision = catalogue_revision(path) if path.exists() else None
        if expected_revision is not None and original_revision != expected_revision:
            raise CatalogueError("The catalogue changed outside this app. Reopen the folder before saving.")
        original_mode = path.stat().st_mode if path.exists() else None
        with tempfile.NamedTemporaryFile(
            prefix=f".{path.stem}-", suffix=path.suffix, dir=path.parent, delete=False
        ) as temporary_file:
            temporary_path = Path(temporary_file.name)

        saved_positions = []
        if path.suffix.lower() == ".xlsx":
            saved_positions = _save_xlsx(path, temporary_path, fieldnames, rows)
        else:
            _save_csv(temporary_path, fieldnames, rows)
        # Windows cannot sync a read-only handle. Both format writers have closed
        # their handles before we sync and replace the file.
        with temporary_path.open("r+b") as temporary_file:
            os.fsync(temporary_file.fileno())
        if original_mode is not None:
            temporary_path.chmod(original_mode & 0o777)
        saved_revision = catalogue_revision(temporary_path)
        current_revision = catalogue_revision(path) if path.exists() else None
        if current_revision != original_revision:
            raise CatalogueError("The catalogue changed outside this app. Reopen the folder before saving.")
        os.replace(temporary_path, path)
        temporary_path = None

        # Update row positions only after replacement succeeds, so another discard
        # still targets the correct worksheet row and failed saves change nothing.
        for row, row_number in saved_positions:
            row[XLSX_ROW_KEY] = row_number
        return saved_revision
    except PermissionError as ex:
        raise CatalogueError(
            f"Could not save {path.name}. Close it in Excel or Numbers and check that the file "
            "and folder are writable, then try Save again. Your edits are still available."
        ) from ex
    except OSError as ex:
        if ex.errno in (errno.ENOSPC, getattr(errno, "EDQUOT", errno.ENOSPC)):
            message = "There is not enough storage space. Free some space, then try Save again."
        else:
            message = f"{ex}. Check that the folder is available, then try Save again."
        raise CatalogueError(f"Could not save {path.name}: {message}") from ex
    except (UnicodeError, csv.Error, BadZipFile, ParseError, ValueError, KeyError, EOFError) as ex:
        raise CatalogueError(f"Could not save {path.name}: {ex}") from ex
    finally:
        if temporary_path is not None:
            try:
                # The original's read-only attribute may have been copied before
                # a failed Windows replacement. Only change the temporary file.
                temporary_path.chmod(0o600)
                temporary_path.unlink(missing_ok=True)
            except OSError:
                # Keep the original save error if cleanup is also denied.
                pass


def _load_csv(path: Path) -> tuple[list[str], list[CatalogueRow]]:
    with path.open(newline="", encoding="utf-8-sig") as f:
        reader = csv.DictReader(f, strict=True)
        fieldnames = list(reader.fieldnames or [])
        _validate_headers(fieldnames)
        rows: list[CatalogueRow] = []
        for row in reader:
            if None in row:
                raise CatalogueError(f"CSV row {reader.line_num} has more values than the header.")
            row = {name: value or "" for name, value in row.items()}
            if REVIEWED_FIELD in row:
                row[REVIEWED_FIELD] = reviewed_text(row[REVIEWED_FIELD], f"CSV row {reader.line_num}")
            rows.append(row)
    return fieldnames, rows


def _save_csv(path: Path, fieldnames: list[str], rows: list[CatalogueRow]) -> None:
    with path.open("w", newline="", encoding="utf-8-sig") as f:
        writer = csv.DictWriter(f, fieldnames=fieldnames, extrasaction="ignore")
        writer.writeheader()
        writer.writerows(rows)


def _openpyxl():
    # A readable error also helps users with an incomplete installation.
    try:
        import openpyxl
    except ImportError as ex:
        raise CatalogueError(
            "Excel support is missing from this installation. Run uv sync to install the app dependencies."
        ) from ex
    return openpyxl


def _cell_text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value)


def reviewed_text(value: object, location: str) -> str:
    """Accept spreadsheet booleans without treating an unknown value as false."""
    text = _cell_text(value).strip().casefold()
    if text in ("true", "1"):
        return REVIEWED_TRUE
    if text in ("", "false", "0"):
        return REVIEWED_FALSE
    raise CatalogueError(f"Invalid Reviewed value in {location}: {value!r}. Use TRUE or FALSE.")


def _xlsx_cell_text(name: str, cell) -> str:
    if name == IMAGE_LINK_FIELD and cell.hyperlink and cell.hyperlink.target:
        # The cell shows a file name; the hyperlink holds the actual crop path.
        return cell.hyperlink.target
    if name == REVIEWED_FIELD:
        return reviewed_text(cell.value, f"cell {cell.coordinate}")
    return _cell_text(cell.value)


def _catalogue_worksheet(wb):
    """Opening another tab in Excel must not hide the Catalogue worksheet."""
    ws = next((sheet for sheet in wb.worksheets if sheet.title.casefold() == "catalogue"), wb.active)
    if ws is None or not hasattr(ws, "iter_rows"):
        raise CatalogueError("The catalogue needs an active worksheet containing records.")
    return ws


def _check_row_deletion(wb, ws, deleted_positions: list[int]) -> None:
    """Refuse deletions that openpyxl cannot safely update on its own."""
    if not deleted_positions:
        return
    dependent_ranges = wb.defined_names or wb.chartsheets or any(
        sheet.tables or sheet.merged_cells or sheet.auto_filter.ref or sheet.data_validations.count
        or sheet.conditional_formatting or sheet.defined_names or sheet.print_area
        or sheet.print_title_rows or sheet.print_title_cols or sheet._charts or sheet._pivots
        for sheet in wb.worksheets
    )
    deleted = set(deleted_positions)
    retained_formulas = any(
        cell.data_type == "f"
        and not (sheet is ws and cell.row in deleted)
        and not (
            sheet is ws and ws.cell(1, cell.column).value == IMAGE_LINK_FIELD
            and isinstance(cell.value, str) and HYPERLINK_FORMULA.fullmatch(cell.value)
        )
        for sheet in wb.worksheets for row in sheet.iter_rows() for cell in row
    )
    if dependent_ranges or retained_formulas:
        raise CatalogueError(
            "Removing a record could change formulas or spreadsheet ranges in this workbook. "
            "Remove the row in Excel or Numbers, then reopen the folder. The catalogue has not been changed."
        )


def _check_xlsx_headers(ws) -> None:
    for column, header in enumerate(ws[1], start=1):
        if _cell_text(header.value).strip():
            continue
        for cells in ws.iter_rows(min_row=2, min_col=column, max_col=column):
            if cells[0].value is not None:
                raise CatalogueError("Every populated column in the catalogue needs a header.")


def _load_xlsx(path: Path) -> tuple[list[str], list[CatalogueRow]]:
    openpyxl = _openpyxl()
    try:
        wb = openpyxl.load_workbook(path)
    except openpyxl.utils.exceptions.InvalidFileException as ex:
        raise CatalogueError(f"Could not read {path.name}: {ex}") from ex
    try:
        ws = _catalogue_worksheet(wb)
        _check_xlsx_headers(ws)
        header = [c.value for c in ws[1]]
        while header and header[-1] is None:
            header.pop()
        fieldnames = [_cell_text(value) for value in header]
        _validate_headers(fieldnames)

        rows: list[CatalogueRow] = []
        for cells in ws.iter_rows(min_row=2, max_col=len(fieldnames)):
            if all(c.value is None for c in cells):
                continue
            row: CatalogueRow = {}
            numeric: set[str] = set()
            for name, cell in zip(fieldnames, cells):
                row[name] = _xlsx_cell_text(name, cell)
                if name == IMAGE_LINK_FIELD and cell.hyperlink and cell.hyperlink.target:
                    row[XLSX_IMAGE_URL_KEY] = True
                if isinstance(cell.value, (int, float)) and not isinstance(cell.value, bool):
                    numeric.add(name)
            row[XLSX_NUMERIC_KEY] = numeric
            row[XLSX_ROW_KEY] = cells[0].row
            rows.append(row)
        return fieldnames, rows
    finally:
        wb.close()


def _xlsx_value(name: str, text: str, numeric: set[str]) -> object:
    """Turn a GUI string back into a cell value of the type the CLI wrote."""
    if text == "":
        return None
    if name == REVIEWED_FIELD and text in (REVIEWED_TRUE, REVIEWED_FALSE):
        return text == REVIEWED_TRUE
    if name in numeric:
        for convert in (int, float):
            try:
                return convert(text)
            except ValueError:
                pass
    return text


def _save_xlsx(
    path: Path, temporary_path: Path, fieldnames: list[str], rows: list[CatalogueRow]
) -> list[tuple[CatalogueRow, int]]:
    openpyxl = _openpyxl()
    from openpyxl.styles import Alignment, Font
    from openpyxl.utils import get_column_letter

    # Reload from disk and leave unchanged cells alone, preserving their types,
    # formulas, comments, formatting, and the rest of the workbook.
    wb = openpyxl.load_workbook(path)
    try:
        ws = _catalogue_worksheet(wb)
        if ws.protection.sheet:
            raise CatalogueError(
                "The Catalogue worksheet is protected. Remove its protection in your spreadsheet app before saving."
            )
        _check_xlsx_headers(ws)
        for col, name in enumerate(fieldnames, start=1):
            header = ws.cell(1, col)
            if header.value is None:
                header.value = name
                if col > 1:
                    header._style = copy(ws.cell(1, col - 1)._style)
                ws.column_dimensions[get_column_letter(col)].width = max(len(name) + 4, 14)

        data_positions = [
            cells[0].row
            for cells in ws.iter_rows(min_row=2, max_col=len(fieldnames))
            if any(cell.value is not None for cell in cells)
        ]
        # Loaded rows carry their worksheet positions, including any gaps caused
        # by blank rows. Plain row dictionaries still work in their existing order.
        has_positions = any(XLSX_ROW_KEY in row for row in rows)
        positions = []
        for index, row in enumerate(rows):
            if has_positions:
                positions.append(row.get(XLSX_ROW_KEY))
            elif index < len(data_positions):
                positions.append(data_positions[index])
            else:
                positions.append(None)
        retained_positions = {position for position in positions if position is not None}
        deleted_positions = [position for position in data_positions if position not in retained_positions]
        _check_row_deletion(wb, ws, deleted_positions)
        for position in reversed(deleted_positions):
            ws.delete_rows(position)

        saved_positions: list[tuple[CatalogueRow, int]] = []
        for row, old_position in zip(rows, positions):
            row_number = (
                old_position - bisect_left(deleted_positions, old_position)
                if old_position is not None
                else max(ws.max_row + 1, 2)
            )
            numeric = row.get(XLSX_NUMERIC_KEY) or set()
            for col, name in enumerate(fieldnames, start=1):
                cell = ws.cell(row_number, col)
                if cell.hyperlink:
                    # Moving a worksheet row does not update its hyperlink refs.
                    cell.hyperlink.ref = cell.coordinate
                text = _cell_text(row.get(name))
                if _xlsx_cell_text(name, cell) == text and not (name == REVIEWED_FIELD and cell.value is None):
                    continue
                if cell.value is None:
                    cell.alignment = Alignment(wrap_text=True, vertical="top")
                if name == IMAGE_LINK_FIELD and text:
                    cell.value = text.replace("\\", "/").rsplit("/", 1)[-1]
                    cell.data_type = "s"
                    cell.hyperlink = text
                    cell.font = Font(color="FF0000FF", underline="single")
                else:
                    cell.value = _xlsx_value(name, text, numeric)
                    if isinstance(cell.value, str):
                        # Transcriptions beginning with '=' are literal text.
                        cell.data_type = "s"
                    if name == IMAGE_LINK_FIELD:
                        cell.hyperlink = None
            saved_positions.append((row, row_number))
        wb.save(temporary_path)
        return saved_positions
    except openpyxl.utils.exceptions.IllegalCharacterError as ex:
        raise CatalogueError(
            "An edited cell contains a control character that Excel cannot store. "
            "Remove the character and try Save again."
        ) from ex
    finally:
        wb.close()


async def page_main(page: ft.Page) -> None:
    page.title = "UMA BCT Review"
    page.window.width = 1000
    page.window.height = 650
    page.window.min_width = 640
    page.window.min_height = 480
    page.window.prevent_close = True
    page.padding = PAGE_PADDING
    page.theme_mode = ft.ThemeMode.SYSTEM

    # Review state. Nothing is loaded until the user picks a folder, so all of
    # this starts empty and is (re)populated by load_folder().
    outputs_dir: Path | None = None
    catalogue_path: Path | None = None
    loaded_revision: str | None = None
    fieldnames: list[str] = []
    rows: list[dict[str, object]] = []
    form_fields: list[str] = []
    current_index = 0
    text_fields: dict[str, ft.TextField] = {}
    focused_text_field: ft.TextField | None = None
    modal_dialogs: list[ft.AlertDialog] = []
    flag_menu_open = False
    folder_picker_open = False
    catalogue_opening = False
    # Flet otherwise treats newly built controls with identical contents as the
    # old controls, leaving Python references detached after a record/folder change.
    control_keys = count()

    file_picker = ft.FilePicker()
    preferences = ft.SharedPreferences()

    def on_field_focus(e: ft.Event) -> None:
        nonlocal focused_text_field
        focused_text_field = e.control

    def on_field_blur(e: ft.Event) -> None:
        nonlocal focused_text_field
        if focused_text_field is e.control:
            focused_text_field = None

    def on_flag_menu_open(e: ft.Event) -> None:
        nonlocal flag_menu_open
        flag_menu_open = True

    def on_flag_menu_close(e: ft.Event) -> None:
        nonlocal flag_menu_open
        flag_menu_open = False

    def update_theme_button() -> None:
        dark = page.theme_mode == ft.ThemeMode.DARK or (
            page.theme_mode == ft.ThemeMode.SYSTEM and page.platform_brightness == ft.Brightness.DARK
        )
        theme_btn.icon = ft.Icons.LIGHT_MODE if dark else ft.Icons.DARK_MODE
        theme_btn.tooltip = "Switch to light mode" if dark else "Switch to dark mode"
        theme_control.label = theme_btn.tooltip
        page.window.brightness = ft.Brightness.DARK if dark else ft.Brightness.LIGHT

    async def toggle_theme(e: ft.Event) -> None:
        if theme_btn.disabled:
            return
        dark = page.theme_mode == ft.ThemeMode.DARK or (
            page.theme_mode == ft.ThemeMode.SYSTEM and page.platform_brightness == ft.Brightness.DARK
        )
        page.theme_mode = ft.ThemeMode.LIGHT if dark else ft.ThemeMode.DARK
        theme_btn.disabled = True
        update_theme_button()
        page.update()
        try:
            saved = await asyncio.wait_for(preferences.set(THEME_PREFERENCE, page.theme_mode.value), timeout=2)
            if not saved:
                raise RuntimeError("Theme preference was not saved")
        except Exception:
            logger.warning("Could not save theme preference", exc_info=True)
            page.show_dialog(ft.SnackBar(content=ft.Text("Theme changed for this session; it couldn't be remembered.")))
        finally:
            theme_btn.disabled = False
            page.update()

    theme_btn = ft.IconButton(on_click=toggle_theme)
    theme_control = ft.Semantics(content=theme_btn, button=True, on_tap=toggle_theme, exclude_semantics=True)

    def on_brightness_change(e: ft.Event) -> None:
        update_theme_button()
        page.update()

    page.on_platform_brightness_change = on_brightness_change

    # Populated by load_folder() with one TextField per catalogue column. Fixed
    # width rather than expand=1, so it sizes to its fields and leaves the rest
    # of the row to the image.
    form_column = ft.Column(
        width=FORM_WIDTH,
        spacing=12,
        scroll=ft.ScrollMode.AUTO,
        horizontal_alignment=ft.CrossAxisAlignment.STRETCH,
    )

    # One color-coded, delete-able chip per flag. Rebuilt on every refresh()
    # from the current row's semicolon-separated flags text. wrap=True lets
    # chips flow onto multiple lines instead of overflowing the fixed-width form.
    flags_row = ft.Row(wrap=True, spacing=8, run_spacing=8)

    def make_flag_chip(flag: str) -> ft.Chip:
        bgcolor, text_color = FLAG_COLORS.get(flag, DEFAULT_FLAG_COLOR)
        chip = ft.Chip(
            label=ft.Text(flag, color=text_color, size=13),
            bgcolor=bgcolor,
            delete_icon_color=text_color,
        )
        chip.data = flag

        def on_delete(e: ft.Event) -> None:
            flags_row.controls.remove(chip)
            refresh_add_flag_menu()
            page.update()

        chip.on_delete = on_delete
        return chip

    def make_add_flag_item(flag: str) -> ft.PopupMenuItem:
        bgcolor, _ = FLAG_COLORS.get(flag, DEFAULT_FLAG_COLOR)
        swatch = ft.Container(width=12, height=12, bgcolor=bgcolor, border_radius=3)

        def pick(e: ft.Event) -> None:
            on_flag_menu_close(e)
            flags_row.controls.append(make_flag_chip(flag))
            refresh_add_flag_menu()
            page.update()

        return ft.PopupMenuItem(content=ft.Row([swatch, ft.Text(flag)], spacing=8), on_click=pick)

    add_flag_btn = ft.PopupMenuButton(
        content="+ Add flag", on_open=on_flag_menu_open,
        on_cancel=on_flag_menu_close, on_select=on_flag_menu_close,
    )

    def refresh_add_flag_menu() -> None:
        """Offer only flags not already applied to this record."""
        existing = {chip.data for chip in flags_row.controls}
        available = [f for f in ALL_FLAGS if f not in existing]
        add_flag_btn.items = [make_add_flag_item(f) for f in available]
        add_flag_btn.disabled = not available

    def set_flags(value: str) -> None:
        flags_row.controls = [make_flag_chip(p) for p in parse_flags(value)]
        refresh_add_flag_menu()

    def get_flags_text() -> str:
        return FLAGS_SEPARATOR.join(chip.data for chip in flags_row.controls)

    flags_block = ft.Column(
        controls=[
            ft.Text("Review Flags", size=12, color=ft.Colors.ON_SURFACE_VARIANT),
            flags_row,
            add_flag_btn,
        ],
        spacing=6,
    )

    reviewed_checkbox = ft.Checkbox(label="Mark as reviewed")

    image_area = ft.Container(expand=True)
    image_viewer: ft.InteractiveViewer | None = None

    async def zoom_in(e: ft.Event) -> None:
        if image_viewer:
            await image_viewer.zoom(1.25)

    async def zoom_out(e: ft.Event) -> None:
        if image_viewer:
            await image_viewer.zoom(0.8)

    async def reset_image(e: ft.Event) -> None:
        if image_viewer:
            await image_viewer.reset()

    zoom_controls = ft.Row(
        controls=[
            ft.TextButton("Zoom in", on_click=zoom_in),
            ft.TextButton("Zoom out", on_click=zoom_out),
            ft.TextButton("Reset / Fit", on_click=reset_image),
            ft.Text("Scroll to zoom; drag to pan", size=12, color=ft.Colors.ON_SURFACE_VARIANT),
        ],
        wrap=True,
        visible=False,
    )

    status_text = ft.Text(size=16, weight=ft.FontWeight.BOLD)
    catalogue_text = ft.Text(size=12, color=ft.Colors.ON_SURFACE_VARIANT, visible=False)

    open_folder_btn = action_button(
        "Open outputs folder",
        ft.Icons.FOLDER_OPEN,
        ft.Colors.PRIMARY_CONTAINER,
        ft.Colors.ON_PRIMARY_CONTAINER,
    )
    open_catalogue_btn = action_button(
        "Open catalogue", ft.Icons.OPEN_IN_NEW,
        ft.Colors.PRIMARY_CONTAINER, ft.Colors.ON_PRIMARY_CONTAINER,
    )
    discard_changes_btn = action_button(
        "Discard changes",
        ft.Icons.CLOSE,
        ft.Colors.ERROR_CONTAINER,
        ft.Colors.ON_ERROR_CONTAINER,
    )
    discard_box_btn = action_button(
        "Discard box",
        ft.Icons.DELETE,
        ft.Colors.ERROR,
        ft.Colors.ON_ERROR,
    )
    save_btn = action_button(
        "Save",
        ft.Icons.SAVE,
        ft.Colors.PRIMARY,
        ft.Colors.ON_PRIMARY,
    )
    prev_btn = action_button(
        "Prev",
        ft.Icons.ARROW_BACK,
        ft.Colors.PRIMARY_CONTAINER,
        ft.Colors.ON_PRIMARY_CONTAINER,
    )
    next_btn = action_button(
        "Next",
        ft.Icons.ARROW_FORWARD,
        ft.Colors.PRIMARY_CONTAINER,
        ft.Colors.ON_PRIMARY_CONTAINER,
    )
    prev_btn.tooltip = "Previous record (Left arrow, outside text fields)"
    next_btn.tooltip = "Next record (Right arrow, outside text fields)"
    view_records_btn = ft.TextButton("View records", on_click=lambda e: go_to(0), visible=False)

    def set_controls_enabled(enabled: bool) -> None:
        for tf in text_fields.values():
            tf.disabled = not enabled
        reviewed_checkbox.disabled = not enabled
        if not enabled:
            # When enabling, availability is governed by refresh_add_flag_menu()
            # instead (disabled once every flag has already been applied).
            add_flag_btn.disabled = True
        discard_changes_btn.disabled = not enabled
        discard_box_btn.disabled = not enabled
        save_btn.disabled = not enabled

    def ask(title: str, message: str, options: list[tuple[str, Callable[[], None] | None]]) -> None:
        """Show a modal dialog with one button per (label, callback) option; the
        dialog closes first, then the chosen callback (if any) runs."""
        handled = False

        def make_handler(callback: Callable[[], None] | None):
            def handler(e: ft.Event) -> None:
                nonlocal handled
                if handled or not dialog.open:
                    return
                handled = True
                # A delayed notification may be above this dialog. Close this
                # instance, and ignore clicks queued during its closing animation.
                dialog.open = False
                page.update(dialog)
                if callback:
                    callback()

            return handler

        dialog = ft.AlertDialog(
            modal=True,
            title=ft.Text(title),
            content=ft.Text(message),
            actions=[ft.TextButton(label, on_click=make_handler(cb)) for label, cb in options],
        )
        modal_dialogs[:] = [item for item in modal_dialogs if item.open]
        modal_dialogs.append(dialog)
        page.show_dialog(dialog)

    def show_error(title: str, message: str) -> None:
        ask(title, message, [("OK", None)])

    def confirm_unsaved_changes(action: str, callback: Callable[[], None]) -> None:
        if not has_unsaved_changes():
            callback()
            return

        def save_then_continue() -> None:
            if save_current():
                callback()

        ask(
            "Unsaved changes",
            f"Save your changes before {action}?",
            [("Cancel", None), ("Don't save", callback), ("Save", save_then_continue)],
        )

    def close_window() -> None:
        page.run_task(page.window.destroy)

    def on_window_event(e: ft.WindowEvent) -> None:
        if e.type == ft.WindowEventType.CLOSE:
            confirm_unsaved_changes("closing the window", close_window)

    page.window.on_event = on_window_event

    def find_next_unreviewed(after_index: int) -> int | None:
        """Search forward from after_index, wrapping around, for an unreviewed
        row. Passing -1 searches the whole list starting from the first row."""
        n = len(rows)
        for offset in range(1, n + 1):
            idx = (after_index + offset) % n
            if rows[idx].get(REVIEWED_FIELD, REVIEWED_FALSE) != REVIEWED_TRUE:
                return idx
        return None

    def update_status() -> None:
        reviewed = sum(row.get(REVIEWED_FIELD) == REVIEWED_TRUE for row in rows)
        catalogue_text.value = f"Editing {catalogue_path.name}" if catalogue_path else ""
        catalogue_text.tooltip = str(catalogue_path) if catalogue_path else None
        catalogue_text.visible = catalogue_path is not None
        open_catalogue_btn.disabled = catalogue_path is None or catalogue_opening
        open_catalogue_btn.tooltip = (
            f"Open {catalogue_path.name} in your default spreadsheet app"
            if catalogue_path else "Open the selected catalogue in your default spreadsheet app"
        )
        if catalogue_path is None:
            status_text.value = "Open an outputs folder to begin"
        elif not rows:
            status_text.value = "This catalogue has no records"
        elif current_index >= len(rows):
            status_text.value = f"All {len(rows)} records reviewed"
        else:
            status_text.value = (
                f"Record {current_index + 1} of {len(rows)} · "
                f"{reviewed} reviewed, {len(rows) - reviewed} remaining"
            )

    def refresh(update: bool = True) -> None:
        """Show the current record, or explain why there is nothing to review."""
        nonlocal image_viewer
        update_status()
        view_records_btn.visible = bool(rows) and current_index >= len(rows)
        image_viewer = None
        zoom_controls.visible = False
        # In the "nothing to review" state current_index == len(rows), so Prev
        # still lets the user step back into the last record.
        prev_btn.disabled = current_index <= 0
        next_btn.disabled = current_index >= len(rows) - 1
        if current_index >= len(rows):
            for tf in text_fields.values():
                tf.value = ""
            flags_row.controls = []
            reviewed_checkbox.value = False
            image_area.content = ft.Text(
                status_text.value,
                text_align=ft.TextAlign.CENTER,
                color=ft.Colors.ON_SURFACE_VARIANT,
            )
            set_controls_enabled(False)
            if update:
                page.update()
            return

        row = rows[current_index]
        for name in form_fields:
            text_fields[name].value = row.get(name, "")
        set_flags(row.get(FLAGS_FIELD, ""))
        reviewed_checkbox.value = row.get(REVIEWED_FIELD, REVIEWED_FALSE) == REVIEWED_TRUE

        image_bytes = None
        image_message = "No image"
        try:
            image_path = resolve_image_path(
                outputs_dir, row.get(IMAGE_LINK_FIELD) or "",
                url_encoded=bool(row.get(XLSX_IMAGE_URL_KEY)),
            )
            if image_path:
                image_bytes = image_path.read_bytes()
                if not image_bytes:
                    image_message = "Couldn't display this image: the file is empty"
        except FileNotFoundError:
            image_message = "Image file not found"
        except (OSError, ValueError):
            image_message = "Couldn't read this image"
        if image_bytes:
            # Read bytes directly rather than pointing Image at the path: the
            # outputs dir is arbitrary and outside Flet's assets_dir.
            image_viewer = ft.InteractiveViewer(
                key=next(control_keys),
                content=ft.Image(
                    src=image_bytes,
                    fit=ft.BoxFit.CONTAIN,
                    expand=True,
                    error_content=ft.Text(
                        "Couldn't display this image: the file may be damaged or unsupported",
                        color=ft.Colors.ON_SURFACE_VARIANT,
                        text_align=ft.TextAlign.CENTER,
                    ),
                ),
                min_scale=1,
                max_scale=8,
                # Lets trackpad scrolling zoom, not just pan, matching mouse-wheel zoom.
                trackpad_scroll_causes_scale=True,
            )
            image_area.content = image_viewer
            zoom_controls.visible = True
        else:
            image_area.content = ft.Text(
                image_message,
                text_align=ft.TextAlign.CENTER,
                color=ft.Colors.ON_SURFACE_VARIANT,
            )

        set_controls_enabled(True)
        if update:
            page.update()

    def load_folder(dir_path: Path) -> None:
        nonlocal outputs_dir, catalogue_path, fieldnames, rows, form_fields, current_index
        nonlocal loaded_revision, focused_text_field

        try:
            candidate = find_catalogue(dir_path)
            if candidate is None:
                raise CatalogueError(f"No catalogue.csv or catalogue.xlsx found in {dir_path}")
            revision = catalogue_revision(candidate)
            loaded_fieldnames, loaded_rows = load_rows(candidate)
            if catalogue_revision(candidate) != revision:
                raise CatalogueError("The catalogue changed while opening it. Please try again.")
        except (CatalogueError, OSError) as ex:
            show_error("Couldn't open folder", str(ex))
            return

        outputs_dir = dir_path
        catalogue_path = candidate
        loaded_revision = revision
        fieldnames, rows = loaded_fieldnames, loaded_rows
        focused_text_field = None
        for name in (FLAGS_FIELD, REVIEWED_FIELD):
            if name not in fieldnames:
                fieldnames.append(name)
        for row in rows:
            row.setdefault(FLAGS_FIELD, "")
            row.setdefault(REVIEWED_FIELD, REVIEWED_FALSE)
        form_fields = [name for name in fieldnames if name not in CUSTOM_FIELDS]
        first_unreviewed = find_next_unreviewed(-1)
        current_index = first_unreviewed if first_unreviewed is not None else len(rows)

        # Rebuild the field inputs from scratch: a different folder's catalogue
        # could have different columns.
        text_fields.clear()
        for name in form_fields:
            text_fields[name] = ft.TextField(
                key=next(control_keys),
                label=name,
                multiline=name in MULTILINE_FIELDS,
                min_lines=3 if name in MULTILINE_FIELDS else 1,
                max_lines=6 if name in MULTILINE_FIELDS else 1,
                read_only=name in READ_ONLY_FIELDS,
                on_focus=on_field_focus,
                on_blur=on_field_blur,
            )
        form_column.controls = [*text_fields.values(), flags_block, reviewed_checkbox]
        if text_fields:
            # The first floating label needs space inside the scroll viewport.
            form_column.controls[0] = ft.Container(
                content=form_column.controls[0], padding=ft.Padding.only(top=8)
            )

        page.title = f"UMA BCT Review — {outputs_dir.name}"
        refresh()

    async def on_open_folder(e: ft.Event) -> None:
        nonlocal folder_picker_open
        if folder_picker_open:
            return
        folder_picker_open = True
        try:
            selected = await file_picker.get_directory_path(
                dialog_title="Select outputs folder"
            )
        except Exception as ex:
            show_error("Couldn't open folder picker", str(ex))
            return
        finally:
            folder_picker_open = False
        if selected:
            confirm_unsaved_changes("opening another folder", lambda: load_folder(Path(selected)))

    def has_unsaved_changes() -> bool:
        if current_index >= len(rows):
            return False
        row = rows[current_index]
        if any((row.get(name) or "") != (text_fields[name].value or "") for name in form_fields):
            return True
        if get_flags_text() != FLAGS_SEPARATOR.join(parse_flags(row.get(FLAGS_FIELD))):
            return True
        saved_reviewed = row.get(REVIEWED_FIELD, REVIEWED_FALSE) == REVIEWED_TRUE
        return bool(reviewed_checkbox.value) != saved_reviewed

    def write_changes(updated_rows: list[dict[str, object]]) -> bool:
        nonlocal rows, loaded_revision
        try:
            revision = save_rows(
                catalogue_path, fieldnames, updated_rows, expected_revision=loaded_revision
            )
        except CatalogueError as ex:
            show_error("Couldn't save catalogue", str(ex))
            return False
        rows = updated_rows
        loaded_revision = revision
        return True

    def save_current() -> bool:
        row = rows[current_index].copy()
        for name in form_fields:
            row[name] = text_fields[name].value or ""
        row[FLAGS_FIELD] = get_flags_text()
        row[REVIEWED_FIELD] = REVIEWED_TRUE if reviewed_checkbox.value else REVIEWED_FALSE
        updated_rows = list(rows)
        updated_rows[current_index] = row
        return write_changes(updated_rows)

    def go_to(index: int) -> None:
        nonlocal current_index
        current_index = index
        refresh()

    def on_save(e: ft.Event) -> None:
        if save_current():
            update_status()
            page.update()
            page.show_dialog(ft.SnackBar(content=ft.Text("Saved"), duration=1500))

    async def launch_catalogue(path: Path) -> None:
        nonlocal catalogue_opening
        try:
            if not path.is_file():
                show_error("Couldn't open catalogue", f"{path.name} could not be found. Check the folder and try again.")
                return
            await asyncio.wait_for(open_in_default_app(path), timeout=10)
        except Exception:
            logger.warning("Could not open catalogue %s", path, exc_info=True)
            show_error(
                "Couldn't open catalogue",
                f"Could not open {path.name}. Check that the folder is available and a default app for {path.suffix} files is set, then try again.",
            )
        finally:
            catalogue_opening = False
            update_status()
            page.update()

    def on_open_catalogue(e: ft.Event) -> None:
        if catalogue_path is None or catalogue_opening or any(dialog.open for dialog in modal_dialogs):
            return
        path = catalogue_path

        def open_saved() -> None:
            nonlocal catalogue_opening
            catalogue_opening = True
            update_status()
            page.update()
            page.run_task(launch_catalogue, path)

        def save_and_open() -> None:
            if save_current():
                open_saved()

        if has_unsaved_changes():
            ask(
                "Unsaved changes",
                f"Save your edits before opening {path.name}? Opening the saved version keeps your unsaved edits in this window.",
                [("Cancel", None), ("Open saved version", open_saved), ("Save and open", save_and_open)],
            )
        else:
            open_saved()

    def on_discard_changes(e: ft.Event) -> None:
        ask(
            "Discard changes?",
            "Any unsaved edits to this box will be lost. Are you sure?",
            [("Cancel", None), ("Discard changes", refresh)],
        )

    def discard_box() -> None:
        # This crop wasn't a real record: drop the row from the catalogue entirely.
        nonlocal current_index
        updated_rows = rows[:current_index] + rows[current_index + 1:]
        if not write_changes(updated_rows):
            return
        current_index = min(current_index, max(len(rows) - 1, 0))
        refresh()

    def on_discard_box(e: ft.Event) -> None:
        ask(
            "Discard box?",
            "This permanently removes the box from the catalogue file. Are you sure?",
            [("Cancel", None), ("Discard box", discard_box)],
        )

    def navigate(step: int) -> None:
        target = current_index + step
        if not 0 <= target < len(rows):
            return
        confirm_unsaved_changes(
            f"going to the {'next' if step > 0 else 'previous'} box", lambda: go_to(target)
        )

    def on_keyboard(e: ft.KeyboardEvent) -> None:
        step = {"Arrow Left": -1, "Arrow Right": 1}.get(e.key)
        if step is None or e.shift or e.ctrl or e.alt or e.meta:
            return
        if focused_text_field is not None or flag_menu_open or folder_picker_open or catalogue_opening:
            return
        if any(dialog.open for dialog in modal_dialogs):
            return
        navigate(step)

    open_folder_btn.on_click = on_open_folder
    open_catalogue_btn.on_click = on_open_catalogue
    page.on_keyboard_event = on_keyboard
    discard_changes_btn.on_click = on_discard_changes
    discard_box_btn.on_click = on_discard_box
    save_btn.on_click = on_save
    prev_btn.on_click = lambda e: navigate(-1)
    next_btn.on_click = lambda e: navigate(1)

    # Left: editable fields. Right: the crop image.
    image_column = ft.Column(
        controls=[image_area, zoom_controls],
        expand=True,
        horizontal_alignment=ft.CrossAxisAlignment.STRETCH,
    )
    top_bar = ft.Row(
        controls=[
            ft.Row(
                controls=[open_folder_btn, open_catalogue_btn, status_text, view_records_btn],
                spacing=16, wrap=True, expand=True,
            ),
            theme_control,
        ],
        vertical_alignment=ft.CrossAxisAlignment.START,
    )

    buttons_row = ft.Row(
        controls=[discard_changes_btn, discard_box_btn, save_btn, prev_btn, next_btn],
        alignment=ft.MainAxisAlignment.START,
        spacing=12,
        wrap=True,
        run_spacing=8,
    )

    body = ft.Row(controls=[form_column, image_column], spacing=24, run_spacing=24, expand=True)
    content_area = ft.Column(
        controls=[body], expand=True, horizontal_alignment=ft.CrossAxisAlignment.STRETCH
    )

    def update_layout(e: ft.PageResizeEvent | None = None) -> None:
        width = e.width if e else page.width or page.window.width
        compact = width < COMPACT_WIDTH
        # Keep the same parents: reparenting live Flet controls detaches their
        # client methods (including zoom), even if they still appear on screen.
        body.wrap = compact
        body.expand = None if compact else True
        content_area.scroll = ft.ScrollMode.AUTO if compact else None
        form_column.width = width - PAGE_PADDING * 2 if compact else FORM_WIDTH
        form_column.scroll = None if compact else ft.ScrollMode.AUTO
        form_column.tight = compact
        image_column.expand = None if compact else True
        image_column.tight = compact
        image_column.width = width - PAGE_PADDING * 2 if compact else None
        image_area.expand = None if compact else True
        image_area.height = 320 if compact else None
        if e:
            page.update()

    page.on_resize = update_layout
    update_layout()
    update_theme_button()

    page.add(
        ft.Column(
            controls=[ft.Column([top_bar, catalogue_text], spacing=4), content_area, buttons_row],
            expand=True,
            spacing=16,
            # Without this, children only take their natural width and the
            # centered empty-state message ends up hugging the left edge.
            horizontal_alignment=ft.CrossAxisAlignment.STRETCH,
        )
    )

    refresh()
    try:
        theme = await asyncio.wait_for(preferences.get(THEME_PREFERENCE), timeout=2)
        if theme in ("light", "dark"):
            page.theme_mode = ft.ThemeMode(theme)
    except Exception:
        logger.warning("Could not read theme preference; using the system theme", exc_info=True)
    update_theme_button()
    page.update()
    try:
        await asyncio.wait_for(page.window.center(), timeout=3)
    except Exception:
        logger.warning("Could not centre the window", exc_info=True)
    finally:
        # A preference or window-manager failure must never leave the app hidden.
        page.window.visible = True
        page.update()


def main() -> None:
    ft.run(page_main, view=ft.AppView.FLET_APP_HIDDEN)


if __name__ == "__main__":
    main()
