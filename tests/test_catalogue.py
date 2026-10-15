import csv
import errno
import os
import tempfile
import unittest
from datetime import datetime
from pathlib import Path
from unittest.mock import patch

import main

try:
    import openpyxl
    from openpyxl.comments import Comment
    from openpyxl.styles import PatternFill
except ImportError:
    openpyxl = None


class CatalogueTestCase(unittest.TestCase):
    def setUp(self):
        self.directory = tempfile.TemporaryDirectory()
        self.addCleanup(self.directory.cleanup)
        self.folder = Path(self.directory.name)
        self.fields = ["Description", main.IMAGE_LINK_FIELD, main.REVIEWED_FIELD]

    def write_csv(self):
        path = self.folder / "catalogue.csv"
        with path.open("w", newline="", encoding="utf-8") as stream:
            writer = csv.writer(stream)
            writer.writerow(self.fields)
            writer.writerow(["Box, with a comma", "crops/box.png", "False"])
        return path

    def assert_no_temp_files(self, path):
        self.assertEqual(list(self.folder.iterdir()), [path])


class CatalogueTests(CatalogueTestCase):
    def test_discovery_preserves_filename_and_selects_newest_with_csv_tie(self):
        csv_path = self.folder / "Catalogue.CSV"
        xlsx_path = self.folder / "CATALOGUE.xlsx"
        csv_path.touch()
        xlsx_path.touch()
        os.utime(csv_path, ns=(1_000_000_000, 1_000_000_000))
        os.utime(xlsx_path, ns=(2_000_000_000, 2_000_000_000))
        self.assertEqual(main.find_catalogue(self.folder), xlsx_path)
        os.utime(csv_path, ns=(2_000_000_000, 2_000_000_000))
        self.assertEqual(main.find_catalogue(self.folder), csv_path)

    def test_discovery_rejects_ambiguous_casing(self):
        # Mock the directory listing so this also runs on case-insensitive disks.
        paths = [self.folder / "catalogue.csv", self.folder / "Catalogue.CSV"]
        with patch.object(Path, "iterdir", return_value=iter(paths)), patch.object(Path, "is_file", return_value=True):
            with self.assertRaisesRegex(main.CatalogueError, "different casing"):
                main.find_catalogue(self.folder)

    def test_csv_normalizes_reviewed_values_and_rejects_unknown_values(self):
        path = self.folder / "catalogue.csv"
        values = ["TRUE", "false", "1", "0", " True ", ""]
        with path.open("w", newline="", encoding="utf-8") as output:
            writer = csv.writer(output)
            writer.writerow(self.fields)
            writer.writerows(["Record", "", value] for value in values)
        _, rows = main.load_rows(path)
        self.assertEqual([row[main.REVIEWED_FIELD] for row in rows], ["True", "False", "True", "False", "True", "False"])
        path.write_text("Description,Reviewed\nRecord,maybe\n", encoding="utf-8")
        with self.assertRaisesRegex(main.CatalogueError, "CSV row 2"):
            main.load_rows(path)

    def test_csv_saves_bom_and_preserves_literal_image_formula_and_identifiers(self):
        path = self.write_csv()
        fields, rows = main.load_rows(path)
        fields.append("Unit Number")
        rows[0]["Unit Number"] = "0007"
        rows[0]["Description"] = 'Café, "box"\nSecond line'
        rows[0][main.IMAGE_LINK_FIELD] = '=HYPERLINK("Box Images/box.png","box.png")'
        main.save_rows(path, fields, rows)
        self.assertTrue(path.read_bytes().startswith(b"\xef\xbb\xbf"))
        self.assertEqual(main.load_rows(path), (fields, rows))

    def test_sync_handle_is_writable_and_sync_failure_keeps_original(self):
        path = self.write_csv()
        fields, rows = main.load_rows(path)
        sync = os.fsync

        def require_writable_handle(fd):
            os.write(fd, b"")  # Fails on a read-only handle, including on POSIX.
            sync(fd)

        with patch("main.os.fsync", side_effect=require_writable_handle):
            main.save_rows(path, fields, rows)
        original = path.read_bytes()
        rows[0]["Description"] = "Unsaved edit"
        with patch("main.os.fsync", side_effect=OSError(errno.ENOSPC, "Disk full")):
            with self.assertRaisesRegex(main.CatalogueError, "storage space"):
                main.save_rows(path, fields, rows)
        self.assertEqual(path.read_bytes(), original)
        self.assert_no_temp_files(path)

    def test_permission_failure_explains_retry_and_keeps_original(self):
        path = self.write_csv()
        original = path.read_bytes()
        fields, rows = main.load_rows(path)
        with patch("main.os.replace", side_effect=PermissionError("File in use")):
            with self.assertRaisesRegex(main.CatalogueError, "Close it in Excel or Numbers"):
                main.save_rows(path, fields, rows)
        self.assertEqual(path.read_bytes(), original)
        self.assert_no_temp_files(path)

    @unittest.skipUnless(os.name == "nt", "Windows replacement of an open file")
    def test_windows_locked_catalogue_preserves_original_and_can_be_retried(self):
        path = self.write_csv()
        fields, rows = main.load_rows(path)
        revision = main.catalogue_revision(path)
        original = path.read_bytes()
        rows[0]["Description"] = "Retry this correction"
        with path.open("rb"):
            with self.assertRaises(main.CatalogueError):
                main.save_rows(path, fields, rows, expected_revision=revision)
            self.assertEqual(path.read_bytes(), original)
        self.assert_no_temp_files(path)
        main.save_rows(path, fields, rows, expected_revision=revision)
        self.assertEqual(main.load_rows(path)[1][0]["Description"], "Retry this correction")

    def test_csv_round_trip_preserves_columns_and_text(self):
        path = self.write_csv()
        fields, rows = main.load_rows(path)
        rows[0]["Description"] = "First line\nSecond line — café"
        fields.append(main.FLAGS_FIELD)
        rows[0][main.FLAGS_FIELD] = "Missing Box"

        main.save_rows(path, fields, rows)

        saved_fields, saved_rows = main.load_rows(path)
        self.assertEqual(saved_fields, fields)
        self.assertEqual(saved_rows, rows)
        self.assert_no_temp_files(path)

    def test_csv_accepts_utf8_bom(self):
        path = self.write_csv()
        path.write_bytes(b"\xef\xbb\xbf" + path.read_bytes())

        fields, rows = main.load_rows(path)

        self.assertEqual(fields, self.fields)
        self.assertEqual(rows[0]["Description"], "Box, with a comma")

    def test_csv_rejects_invalid_headers(self):
        path = self.folder / "catalogue.csv"
        for content in ("", "Description,Description\nA,B\n", "Description,\nA,B\n"):
            with self.subTest(content=content):
                path.write_text(content, encoding="utf-8")
                with self.assertRaises(main.CatalogueError):
                    main.load_rows(path)

    def test_csv_rejects_an_unclosed_quoted_field(self):
        path = self.folder / "catalogue.csv"
        path.write_text('Description,Reviewed\n"Unclosed description,True\n', encoding="utf-8")
        with self.assertRaises(main.CatalogueError):
            main.load_rows(path)

    def test_csv_write_failure_keeps_original_and_removes_temp_file(self):
        path = self.write_csv()
        original = path.read_bytes()
        fields, rows = main.load_rows(path)
        rows[0]["Description"] = "Changed"

        with patch.object(csv.DictWriter, "writerows", side_effect=OSError("Disk full")):
            with self.assertRaises(main.CatalogueError):
                main.save_rows(path, fields, rows)

        self.assertEqual(path.read_bytes(), original)
        self.assert_no_temp_files(path)

    def test_csv_replace_failure_keeps_original_and_removes_temp_file(self):
        path = self.write_csv()
        original = path.read_bytes()
        fields, rows = main.load_rows(path)

        with patch("main.os.replace", side_effect=OSError("Permission denied")):
            with self.assertRaises(main.CatalogueError):
                main.save_rows(path, fields, rows)

        self.assertEqual(path.read_bytes(), original)
        self.assert_no_temp_files(path)

    def test_csv_refuses_to_overwrite_external_changes(self):
        path = self.write_csv()
        fields, rows = main.load_rows(path)
        revision = main.catalogue_revision(path)
        external_content = path.read_bytes().replace(b"Box, with a comma", b"External edit")
        path.write_bytes(external_content)
        rows[0]["Description"] = "GUI edit"

        with self.assertRaises(main.CatalogueError):
            main.save_rows(path, fields, rows, expected_revision=revision)

        self.assertEqual(path.read_bytes(), external_content)
        self.assert_no_temp_files(path)

    def test_csv_detects_external_changes_during_save(self):
        path = self.write_csv()
        fields, rows = main.load_rows(path)
        revision = main.catalogue_revision(path)
        external_content = path.read_bytes().replace(b"Box, with a comma", b"External edit")
        original_writerows = csv.DictWriter.writerows

        def write_then_change_original(writer, values):
            original_writerows(writer, values)
            path.write_bytes(external_content)

        with patch.object(csv.DictWriter, "writerows", new=write_then_change_original):
            with self.assertRaises(main.CatalogueError):
                main.save_rows(path, fields, rows, expected_revision=revision)

        self.assertEqual(path.read_bytes(), external_content)
        self.assert_no_temp_files(path)


@unittest.skipIf(openpyxl is None, "Install the excel extra to test XLSX catalogues")
class ExcelCatalogueTests(CatalogueTestCase):
    @unittest.skipUnless(os.name == "nt", "Windows replacement of an open workbook")
    def test_windows_locked_workbook_preserves_original_and_can_be_retried(self):
        path = self.write_workbook()
        fields, rows = main.load_rows(path)
        revision = main.catalogue_revision(path)
        original = path.read_bytes()
        rows[0]["Description"] = "Retry this correction"
        with path.open("rb"):
            with self.assertRaises(main.CatalogueError):
                main.save_rows(path, fields, rows, expected_revision=revision)
            self.assertEqual(path.read_bytes(), original)
        self.assert_no_temp_files(path)
        main.save_rows(path, fields, rows, expected_revision=revision)
        self.assertEqual(main.load_rows(path)[1][0]["Description"], "Retry this correction")

    def test_named_catalogue_is_read_and_saved_when_notes_tab_is_active(self):
        path = self.write_workbook()
        workbook = self.read_workbook(path)
        workbook.active.title = "Catalogue"
        workbook.active = 1
        workbook.save(path)
        fields, rows = main.load_rows(path)
        self.assertEqual(rows[0]["Description"], "Original description")
        rows[0]["Description"] = "Correction"
        main.save_rows(path, fields, rows)
        saved = self.read_workbook(path)
        self.assertEqual(saved["Catalogue"]["A2"].value, "Correction")
        self.assertEqual(saved.active.title, "Notes")
        self.assertEqual(saved["Notes"]["A1"].value, "Other sheet stays intact")

    def test_xlsx_string_booleans_are_recognized_and_invalid_values_rejected(self):
        path = self.write_workbook()
        workbook = self.read_workbook(path)
        for value, expected in (("TRUE", "True"), ("0", "False"), (" false ", "False")):
            with self.subTest(value=value):
                workbook.active["C2"] = value
                workbook.save(path)
                self.assertEqual(main.load_rows(path)[1][0][main.REVIEWED_FIELD], expected)
        workbook.active["C2"] = "maybe"
        workbook.save(path)
        with self.assertRaisesRegex(main.CatalogueError, "cell C2"):
            main.load_rows(path)

    def test_xlsx_native_encoded_links_and_formula_links_keep_their_values(self):
        path = self.write_workbook()
        workbook = self.read_workbook(path)
        link = "Box%20Images/box%23one%25.png"
        workbook.active["B2"].hyperlink = link
        formula = '=HYPERLINK("Box Images/100%20.png","Box")'
        workbook.active.append(["Formula image", formula, False])
        workbook.save(path)
        fields, rows = main.load_rows(path)
        self.assertTrue(rows[0][main.XLSX_IMAGE_URL_KEY])
        self.assertNotIn(main.XLSX_IMAGE_URL_KEY, rows[1])
        self.assertEqual(main.resolve_image_path(self.folder, rows[0][main.IMAGE_LINK_FIELD], url_encoded=True), self.folder / "Box Images/box#one%.png")
        main.save_rows(path, fields, rows)
        saved = self.read_workbook(path)
        self.assertEqual(saved.active["B2"].hyperlink.target, link)
        self.assertEqual(saved.active["B3"].value, formula)
        self.assertEqual(saved.active["B3"].data_type, "f")

    def test_xlsx_allows_deleting_rows_with_remaining_literal_image_formulas(self):
        path = self.write_workbook()
        workbook = self.read_workbook(path)
        formula = '=HYPERLINK("Box Images/second.png","Second")'
        workbook.active.append(["Second", formula, False])
        workbook.save(path)
        fields, rows = main.load_rows(path)
        main.save_rows(path, fields, rows[1:])
        self.assertEqual(self.read_workbook(path).active["B2"].value, formula)

    def test_xlsx_refuses_deletions_with_retained_formulas_or_ranges(self):
        for feature in ("formula", "other_sheet_formula", "filter", "table", "merged", "named_range",
                        "other_sheet_name", "other_sheet_chart", "chartsheet"):
            with self.subTest(feature=feature):
                path = self.write_workbook()
                workbook = self.read_workbook(path)
                sheet = workbook.active
                sheet.append(["Second", "second.png", False])
                if feature == "formula":
                    sheet["D3"] = "=D2"
                elif feature == "other_sheet_formula":
                    workbook["Notes"]["A2"] = "=Sheet!D2"
                elif feature == "filter":
                    sheet.auto_filter.ref = "A1:H3"
                elif feature == "table":
                    from openpyxl.worksheet.table import Table
                    sheet.add_table(Table(displayName="Records", ref="A1:C3"))
                elif feature == "merged":
                    sheet.merge_cells("D3:E3")
                elif feature in ("named_range", "other_sheet_name"):
                    from openpyxl.workbook.defined_name import DefinedName
                    names = workbook.defined_names if feature == "named_range" else workbook["Notes"].defined_names
                    names.add(DefinedName("RecordCount", attr_text="Sheet!$D$2"))
                else:
                    from openpyxl.chart import BarChart, Reference
                    chart = BarChart()
                    chart.add_data(Reference(sheet, min_col=4, min_row=2, max_row=3))
                    destination = workbook["Notes"] if feature == "other_sheet_chart" else workbook.create_chartsheet("Chart")
                    destination.add_chart(chart)
                workbook.save(path)
                original = path.read_bytes()
                fields, rows = main.load_rows(path)
                with self.assertRaisesRegex(main.CatalogueError, "formulas or spreadsheet ranges"):
                    main.save_rows(path, fields, rows[1:])
                self.assertEqual(path.read_bytes(), original)
                self.assert_no_temp_files(path)
                # Ordinary corrections are still allowed in these workbooks.
                rows[0]["Description"] = "Correction"
                main.save_rows(path, fields, rows)

    def test_xlsx_protection_and_invalid_control_characters_fail_safely(self):
        path = self.write_workbook()
        workbook = self.read_workbook(path)
        workbook.active.protection.sheet = True
        workbook.save(path)
        fields, rows = main.load_rows(path)
        original = path.read_bytes()
        with self.assertRaisesRegex(main.CatalogueError, "protected"):
            main.save_rows(path, fields, rows)
        self.assertEqual(path.read_bytes(), original)
        workbook.active.protection.sheet = False
        workbook.save(path)
        original = path.read_bytes()
        rows[0]["Description"] = "bad\x01text"
        with self.assertRaisesRegex(main.CatalogueError, "control character"):
            main.save_rows(path, fields, rows)
        self.assertEqual(path.read_bytes(), original)
        self.assert_no_temp_files(path)

    def write_workbook(self):
        path = self.folder / "catalogue.xlsx"
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        self.fields = ["Description", main.IMAGE_LINK_FIELD, main.REVIEWED_FIELD, "Count", "Date", "Selected", "Formula", "Website"]
        sheet.append(self.fields)
        sheet.append([
            "Original description", "box.png", False, 12, datetime(2024, 5, 6),
            True, "=D2*2", "Archive website",
        ])
        sheet["B2"].hyperlink = "crops/box.png"
        sheet["A2"].comment = Comment("Keep this note", "Reviewer")
        sheet["A2"].fill = PatternFill(fill_type="solid", fgColor="FFFFCC00")
        sheet["D2"].number_format = "0.00"
        sheet["E2"].number_format = "yyyy-mm-dd"
        sheet["H2"].hyperlink = "https://example.com/archive"
        workbook.create_sheet("Notes")["A1"] = "Other sheet stays intact"
        workbook.save(path)
        workbook.close()
        return path

    def read_workbook(self, path):
        workbook = openpyxl.load_workbook(path)
        self.addCleanup(workbook.close)
        return workbook

    def test_xlsx_rejects_populated_columns_without_headers(self):
        path = self.folder / "catalogue.xlsx"
        workbook = openpyxl.Workbook()
        workbook.active.append(["Description", main.REVIEWED_FIELD])
        workbook.active.append(["Record", False])
        workbook.active["C3"] = "A note without a column header"
        workbook.save(path)
        workbook.close()
        original = path.read_bytes()

        with self.assertRaises(main.CatalogueError):
            main.load_rows(path)

        self.assertEqual(path.read_bytes(), original)

    def test_xlsx_accepts_empty_formatted_columns_without_headers(self):
        path = self.folder / "catalogue.xlsx"
        workbook = openpyxl.Workbook()
        workbook.active.append(["Description", main.REVIEWED_FIELD])
        workbook.active.append(["Record", False])
        workbook.active["C3"].fill = PatternFill(fill_type="solid", fgColor="FFFFCC00")
        workbook.save(path)
        workbook.close()

        fields, rows = main.load_rows(path)

        self.assertEqual(fields, ["Description", main.REVIEWED_FIELD])
        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0]["Description"], "Record")

    def test_xlsx_save_does_not_replace_populated_unnamed_columns_with_flags(self):
        path = self.folder / "catalogue.xlsx"
        fields = ["Description", main.REVIEWED_FIELD, main.FLAGS_FIELD]
        rows = [{"Description": "Changed", main.REVIEWED_FIELD: "False", main.FLAGS_FIELD: "Missing Box"}]
        for position in ("C2", "C3"):
            with self.subTest(position=position):
                workbook = openpyxl.Workbook()
                workbook.active.append(["Description", main.REVIEWED_FIELD])
                workbook.active.append(["Record", False])
                workbook.active[position] = "Keep this note"
                workbook.save(path)
                workbook.close()
                original = path.read_bytes()

                with self.assertRaises(main.CatalogueError):
                    main.save_rows(path, fields, rows)

                self.assertEqual(path.read_bytes(), original)
                self.assert_no_temp_files(path)

    def test_xlsx_round_trip_preserves_metadata_and_original_types(self):
        path = self.write_workbook()
        fields, rows = main.load_rows(path)
        rows[0]["Description"] = "Corrected description"
        fields.append(main.FLAGS_FIELD)
        rows[0][main.FLAGS_FIELD] = "Missing Box"

        main.save_rows(path, fields, rows)

        workbook = self.read_workbook(path)
        sheet = workbook.active
        self.assertEqual(sheet["A2"].value, "Corrected description")
        self.assertEqual(sheet["A2"].comment.text, "Keep this note")
        self.assertEqual(sheet["A2"].fill.fgColor.rgb, "FFFFCC00")
        self.assertEqual(sheet["B2"].hyperlink.target, "crops/box.png")
        self.assertIs(sheet["C2"].value, False)
        self.assertEqual(sheet["D2"].value, 12)
        self.assertEqual(sheet["D2"].number_format, "0.00")
        self.assertEqual(sheet["E2"].value, datetime(2024, 5, 6))
        self.assertEqual(sheet["E2"].number_format, "yyyy-mm-dd")
        self.assertIs(sheet["F2"].value, True)
        self.assertEqual(sheet["G2"].value, "=D2*2")
        self.assertEqual(sheet["G2"].data_type, "f")
        self.assertEqual(sheet["H2"].hyperlink.target, "https://example.com/archive")
        self.assertEqual(sheet["I1"].value, main.FLAGS_FIELD)
        self.assertEqual(sheet["I2"].value, "Missing Box")
        self.assertEqual(workbook["Notes"]["A1"].value, "Other sheet stays intact")
        self.assert_no_temp_files(path)

    def test_xlsx_preserves_and_saves_literal_equals_text(self):
        path = self.write_workbook()
        workbook = self.read_workbook(path)
        workbook.active["A2"] = "=1+2"
        workbook.active["A2"].data_type = "s"
        workbook.save(path)
        fields, rows = main.load_rows(path)

        main.save_rows(path, fields, rows)
        saved = self.read_workbook(path)
        self.assertEqual(saved.active["A2"].value, "=1+2")
        self.assertEqual(saved.active["A2"].data_type, "s")

        rows[0]["Description"] = "=3+4"
        main.save_rows(path, fields, rows)
        edited = self.read_workbook(path)
        self.assertEqual(edited.active["A2"].value, "=3+4")
        self.assertEqual(edited.active["A2"].data_type, "s")

    def test_xlsx_normalizes_numeric_and_boolean_reviewed_values(self):
        path = self.folder / "catalogue.xlsx"
        workbook = openpyxl.Workbook()
        workbook.active.append(self.fields)
        for value in (True, False, 1, 0):
            workbook.active.append(["Record", "crop.png", value])
        workbook.save(path)
        workbook.close()

        fields, rows = main.load_rows(path)

        self.assertEqual([row[main.REVIEWED_FIELD] for row in rows], ["True", "False", "True", "False"])
        rows[1][main.REVIEWED_FIELD] = "True"
        main.save_rows(path, fields, rows)
        saved = self.read_workbook(path)
        self.assertIs(saved.active["C3"].value, True)
        self.assertEqual(saved.active["C4"].value, 1)
        self.assertEqual(saved.active["C5"].value, 0)

    def test_xlsx_deletes_correct_record_with_blank_row_gaps(self):
        path = self.write_workbook()
        workbook = self.read_workbook(path)
        sheet = workbook.active
        sheet.append([None] * len(self.fields))
        sheet.append(["Keep this record", "second.png", False])
        sheet["A4"].comment = Comment("Keep second note", "Reviewer")
        workbook.save(path)
        fields, rows = main.load_rows(path)
        self.assertEqual(len(rows), 2)

        main.save_rows(path, fields, rows[1:])

        saved = self.read_workbook(path)
        self.assertIsNone(saved.active["A2"].value)
        self.assertEqual(saved.active["A3"].value, "Keep this record")
        self.assertEqual(saved.active["A3"].comment.text, "Keep second note")
        self.assertEqual(len(main.load_rows(path)[1]), 1)

    def test_xlsx_repeated_deletes_keep_correct_row_positions(self):
        path = self.write_workbook()
        workbook = self.read_workbook(path)
        workbook.active.append(["Second record", "second.png", False])
        workbook.active.append(["Third record", "third.png", False])
        workbook.active["A4"].comment = Comment("Third record note", "Reviewer")
        workbook.save(path)
        fields, rows = main.load_rows(path)
        remaining = rows[1:]

        revision = main.save_rows(path, fields, remaining)
        main.save_rows(path, fields, remaining[1:], expected_revision=revision)

        saved = self.read_workbook(path)
        self.assertEqual(saved.active["A2"].value, "Third record")
        self.assertEqual(saved.active["A2"].comment.text, "Third record note")
        self.assertEqual(len(main.load_rows(path)[1]), 1)

    def test_xlsx_delete_moves_hyperlinks_without_creating_ghost_rows(self):
        path = self.write_workbook()
        workbook = self.read_workbook(path)
        sheet = workbook.active
        sheet.append(["Second record", "second.png", False, None, None, None, None, "Second website"])
        sheet["B3"].hyperlink = "crops/second.png"
        sheet["H3"].hyperlink = "https://example.com/second"
        workbook.save(path)
        fields, rows = main.load_rows(path)

        main.save_rows(path, fields, rows[1:])

        saved_fields, saved_rows = main.load_rows(path)
        self.assertEqual(saved_fields, fields)
        self.assertEqual(len(saved_rows), 1)
        self.assertEqual(saved_rows[0]["Description"], "Second record")
        self.assertEqual(saved_rows[0][main.IMAGE_LINK_FIELD], "crops/second.png")
        saved = self.read_workbook(path)
        sheet = saved.active
        self.assertEqual(sheet.max_row, 2)
        self.assertEqual(sheet["B2"].hyperlink.target, "crops/second.png")
        self.assertEqual(sheet["B2"].hyperlink.ref, "B2")
        self.assertEqual(sheet["H2"].hyperlink.target, "https://example.com/second")
        self.assertEqual(sheet["H2"].hyperlink.ref, "H2")

    def test_xlsx_write_failure_keeps_original_and_removes_temp_file(self):
        path = self.write_workbook()
        original = path.read_bytes()
        fields, rows = main.load_rows(path)

        with patch.object(openpyxl.Workbook, "save", side_effect=OSError("Disk full")):
            with self.assertRaises(main.CatalogueError):
                main.save_rows(path, fields, rows)

        self.assertEqual(path.read_bytes(), original)
        self.assert_no_temp_files(path)

    def test_xlsx_refuses_to_overwrite_external_changes(self):
        path = self.write_workbook()
        fields, rows = main.load_rows(path)
        revision = main.catalogue_revision(path)
        workbook = self.read_workbook(path)
        workbook["Notes"]["A1"] = "An external correction"
        workbook.save(path)
        external_content = path.read_bytes()

        with self.assertRaises(main.CatalogueError):
            main.save_rows(path, fields, rows, expected_revision=revision)

        self.assertEqual(path.read_bytes(), external_content)
        self.assert_no_temp_files(path)

    def test_corrupt_xlsx_has_readable_error(self):
        path = self.folder / "catalogue.xlsx"
        path.write_text("This is not an Excel workbook", encoding="utf-8")

        with self.assertRaises(main.CatalogueError):
            main.load_rows(path)


if __name__ == "__main__":
    unittest.main()
