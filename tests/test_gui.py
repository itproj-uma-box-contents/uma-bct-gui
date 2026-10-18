import asyncio
import base64
import csv
import inspect
import os
import tempfile
import unittest
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import AsyncMock, patch

import flet as ft

import main


class TestPage:
    """Keep real Flet controls, without launching a desktop window."""

    def __init__(self):
        self.controls = []
        self.dialogs = []
        self.tasks = []
        self.window = SimpleNamespace(destroy=AsyncMock(), center=AsyncMock(), prevent_close=False)
        self.width = 1000
        self.height = 618
        self.platform_brightness = ft.Brightness.LIGHT

    def add(self, control):
        self.controls.append(control)

    def update(self, *controls):
        self.dialogs[:] = [dialog for dialog in self.dialogs if dialog.open]

    def show_dialog(self, dialog):
        dialog.open = True
        self.dialogs.append(dialog)

    def pop_dialog(self):
        self.dialogs.pop()

    def run_task(self, callback, *args):
        try:
            loop = asyncio.get_running_loop()
        except RuntimeError:
            return asyncio.run(callback(*args))
        task = loop.create_task(callback(*args))
        self.tasks.append(task)
        return task


def walk_controls(control):
    yield control
    for child in getattr(control, "controls", []):
        yield from walk_controls(child)
    content = getattr(control, "content", None)
    if content is not None and not isinstance(content, str):
        yield from walk_controls(content)


class GuiTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.folder = Path(self.temp.name)
        self.path = self.folder / "catalogue.csv"
        self.write_catalogue()
        self.preferences = {}
        self.get_preference = AsyncMock(side_effect=lambda key: self.preferences.get(key))
        self.set_preference = AsyncMock(side_effect=self.remember_preference)
        for method, mock in (("get", self.get_preference), ("set", self.set_preference)):
            patcher = patch.object(ft.SharedPreferences, method, mock)
            patcher.start()
            self.addCleanup(patcher.stop)
        self.launch = AsyncMock()
        patcher = patch.object(main, "open_in_default_app", self.launch)
        patcher.start()
        self.addCleanup(patcher.stop)
        self.page = TestPage()
        asyncio.run(main.page_main(self.page))

    def remember_preference(self, key, value):
        self.preferences[key] = value
        return True

    def write_catalogue(self, reviewed=False, include_flags=True):
        fields = ["Description", main.IMAGE_LINK_FIELD, main.REVIEWED_FIELD]
        if include_flags:
            fields.append(main.FLAGS_FIELD)
        with self.path.open("w", newline="", encoding="utf-8") as output:
            writer = csv.DictWriter(output, fieldnames=fields)
            writer.writeheader()
            for description in ("First box", "Second box"):
                writer.writerow({"Description": description, main.REVIEWED_FIELD: str(reviewed)})

    def controls(self):
        for control in self.page.controls:
            yield from walk_controls(control)

    def button(self, label):
        return next(c for c in self.controls() if getattr(c, "content", None) == label)

    def description(self):
        return next(c for c in self.controls() if isinstance(c, ft.TextField) and getattr(c.label, "value", c.label) == "Description")

    def click(self, control):
        handler = control.on_click
        if inspect.iscoroutinefunction(handler):
            asyncio.run(handler(None))
        else:
            handler(None)

    def choose(self, label):
        dialog = self.page.dialogs[-1]
        self.click(next(button for button in dialog.actions if button.content == label))

    def open_folder(self, folder=None):
        with patch.object(ft.FilePicker, "get_directory_path", new=AsyncMock(return_value=str(folder or self.folder))):
            self.click(self.button("Open outputs folder"))

    def key(self, key, **modifiers):
        values = dict(shift=False, ctrl=False, alt=False, meta=False)
        values.update(modifiers)
        self.page.on_keyboard_event(SimpleNamespace(key=key, **values))

    def test_arrow_navigation_matches_buttons_and_stops_at_boundaries(self):
        self.key("Arrow Left")
        self.key("Arrow Right")
        self.assertEqual(self.page.dialogs, [])
        self.open_folder()
        self.key("Arrow Left")
        self.assertEqual(self.description().value, "First box")
        self.key("Arrow Right")
        self.assertEqual(self.description().value, "Second box")
        self.key("Arrow Right")
        self.assertEqual(self.description().value, "Second box")
        self.key("Arrow Left")
        self.assertEqual(self.description().value, "First box")

    def test_arrows_leave_text_fields_and_modified_keys_alone(self):
        self.open_folder()
        field = self.description()
        field.on_focus(SimpleNamespace(control=field))
        self.key("Arrow Right")
        self.assertEqual(field.value, "First box")
        field.on_blur(SimpleNamespace(control=field))
        for modifier in ("shift", "ctrl", "alt", "meta"):
            self.key("Arrow Right", **{modifier: True})
            self.assertEqual(field.value, "First box")
        self.key("Arrow Up")
        self.assertEqual(field.value, "First box")
        self.key("Arrow Right")
        self.assertEqual(field.value, "Second box")

    def test_late_blur_from_another_field_does_not_enable_arrow_navigation(self):
        self.path.write_text("Description,Other\nFirst box,One\nSecond box,Two\n", encoding="utf-8")
        self.open_folder()
        first = self.description()
        second = next(c for c in self.controls() if isinstance(c, ft.TextField) and getattr(c.label, "value", c.label) == "Other")
        first.on_focus(SimpleNamespace(control=first))
        second.on_focus(SimpleNamespace(control=second))
        first.on_blur(SimpleNamespace(control=first))
        self.key("Arrow Right")
        self.assertEqual(first.value, "First box")
        second.on_blur(SimpleNamespace(control=second))
        self.key("Arrow Right")
        self.assertEqual(first.value, "Second box")

    def test_arrows_use_unsaved_prompt_and_do_not_open_duplicate_confirmations(self):
        self.open_folder()
        self.description().value = "Unsaved correction"
        self.key("Arrow Right")
        confirmation = self.page.dialogs[-1]
        self.key("Arrow Right")
        self.key("Arrow Left")
        self.assertEqual(self.page.dialogs, [confirmation])
        self.assertEqual(self.description().value, "Unsaved correction")
        self.choose("Cancel")
        self.key("Arrow Right")
        self.choose("Save")
        self.assertEqual(self.description().value, "Second box")
        self.assertEqual(main.load_rows(self.path)[1][0]["Description"], "Unsaved correction")

    def test_arrow_navigation_stays_on_record_if_save_fails(self):
        self.open_folder()
        self.description().value = "Unsaved correction"
        self.key("Arrow Right")
        with patch.object(main, "save_rows", side_effect=main.CatalogueError("Disk full")):
            self.choose("Save")
        self.assertEqual(self.description().value, "Unsaved correction")
        self.assertEqual(main.load_rows(self.path)[1][0]["Description"], "First box")

    def test_arrows_do_not_change_the_record_targeted_by_a_discard_dialog(self):
        self.open_folder()
        self.click(self.button("Discard box"))
        self.key("Arrow Right")
        self.choose("Discard box")
        self.assertEqual([row["Description"] for row in main.load_rows(self.path)[1]], ["Second box"])

    def test_queued_button_clicks_do_not_change_the_record_targeted_by_a_modal(self):
        self.open_folder()
        self.click(self.button("Discard box"))
        confirmation = self.page.dialogs[-1]
        for label in ("Next", "Prev", "View records", "Save", "Discard box", "Discard changes", "Open catalogue"):
            self.click(self.button(label))
        with patch.object(ft.FilePicker, "get_directory_path", new=AsyncMock()) as picker:
            self.click(self.button("Open outputs folder"))
            picker.assert_not_awaited()
        self.assertEqual(self.page.dialogs, [confirmation])
        self.assertEqual(self.description().value, "First box")
        self.launch.assert_not_awaited()
        self.choose("Discard box")
        self.assertEqual([row["Description"] for row in main.load_rows(self.path)[1]], ["Second box"])

    def test_repeated_close_events_share_one_unsaved_confirmation(self):
        self.open_folder()
        self.description().value = "Unsaved correction"
        event = SimpleNamespace(type=ft.WindowEventType.CLOSE)
        self.page.window.on_event(event)
        confirmation = self.page.dialogs[-1]
        self.page.window.on_event(event)
        self.page.window.on_event(event)
        self.assertEqual(self.page.dialogs, [confirmation])
        self.choose("Cancel")
        self.page.window.destroy.assert_not_awaited()
        self.assertEqual(self.description().value, "Unsaved correction")

    def test_queued_close_during_discard_confirmation_does_not_close_the_window(self):
        self.open_folder()
        self.click(self.button("Discard box"))
        self.page.window.on_event(SimpleNamespace(type=ft.WindowEventType.CLOSE))
        self.page.window.destroy.assert_not_awaited()
        self.choose("Cancel")
        self.assertEqual(self.description().value, "First box")

    def test_arrows_leave_the_flag_menu_alone(self):
        self.open_folder()
        menu = self.button("+ Add flag")
        for close in (menu.on_cancel, menu.on_select):
            menu.on_open(None)
            self.key("Arrow Right")
            self.assertEqual(self.description().value, "First box")
            close(None)
            self.key("Arrow Right")
            self.assertEqual(self.description().value, "Second box")
            self.key("Arrow Left")

    def test_arrows_are_ignored_while_the_native_folder_picker_is_open(self):
        self.open_folder()

        async def exercise():
            started = asyncio.Event()
            release = asyncio.Event()

            async def picker(**kwargs):
                started.set()
                await release.wait()
                return None

            with patch.object(ft.FilePicker, "get_directory_path", new=AsyncMock(side_effect=picker)):
                task = asyncio.create_task(self.button("Open outputs folder").on_click(None))
                await asyncio.wait_for(started.wait(), timeout=1)
                self.key("Arrow Right")
                self.assertEqual(self.description().value, "First box")
                release.set()
                await task
            self.key("Arrow Right")
            self.assertEqual(self.description().value, "Second box")

        asyncio.run(exercise())

    def test_folder_picker_does_not_replace_a_modal_that_arrives_while_waiting(self):
        self.open_folder()
        other = self.folder / "other"
        other.mkdir()
        (other / "catalogue.csv").write_text("Description\nOther record\n", encoding="utf-8")
        handler = self.button("Open outputs folder").on_click
        show_error = inspect.getclosurevars(handler).nonlocals["show_error"]

        async def picker(**kwargs):
            # Model an asynchronous error arriving while the native picker is open.
            show_error("Delayed error", "Keep this message available")
            return str(other)

        with patch.object(ft.FilePicker, "get_directory_path", new=AsyncMock(side_effect=picker)):
            self.click(self.button("Open outputs folder"))
        self.assertEqual(self.description().value, "First box")
        self.assertEqual(self.page.dialogs[-1].title.value, "Delayed error")
        self.choose("OK")
        self.open_folder(other)
        self.assertEqual(self.description().value, "Other record")

    def test_open_catalogue_is_disabled_until_a_file_is_loaded(self):
        self.assertTrue(self.button("Open catalogue").disabled)
        self.click(self.button("Open catalogue"))
        self.launch.assert_not_awaited()
        self.open_folder()
        self.assertFalse(self.button("Open catalogue").disabled)
        self.click(self.button("Open catalogue"))
        self.launch.assert_awaited_once_with(self.path)

    def test_open_catalogue_handles_special_characters_and_xlsx_selection(self):
        import openpyxl

        folder = self.folder / "Review # café"
        folder.mkdir()
        path = folder / "Catalogue.xlsx"
        workbook = openpyxl.Workbook()
        workbook.active.append(["Description", main.REVIEWED_FIELD])
        workbook.active.append(["Workbook record", False])
        workbook.save(path)
        workbook.close()
        self.open_folder(folder)
        self.click(self.button("Open catalogue"))
        self.launch.assert_awaited_once_with(path)
        self.assertIn("Catalogue.xlsx", self.button("Open catalogue").tooltip)

    def test_open_saved_version_preserves_unsaved_gui_edits(self):
        self.open_folder()
        original = self.path.read_bytes()
        self.description().value = "Unsaved correction"
        self.click(self.button("Open catalogue"))
        self.choose("Cancel")
        self.launch.assert_not_awaited()
        self.click(self.button("Open catalogue"))
        self.choose("Open saved version")
        self.launch.assert_awaited_once()
        self.assertEqual(self.path.read_bytes(), original)
        self.assertEqual(self.description().value, "Unsaved correction")
        self.key("Arrow Right")
        self.assertEqual(self.page.dialogs[-1].title.value, "Unsaved changes")

    def test_save_and_open_writes_before_opening_and_failure_does_not_open(self):
        self.open_folder()
        self.description().value = "Correction"
        self.click(self.button("Open catalogue"))
        with patch.object(main, "save_rows", side_effect=main.CatalogueError("Disk full")):
            self.choose("Save and open")
        self.launch.assert_not_awaited()
        self.assertEqual(self.description().value, "Correction")
        self.choose("OK")

        async def check_saved(*args, **kwargs):
            self.assertEqual(main.load_rows(self.path)[1][0]["Description"], "Correction")

        self.launch.side_effect = check_saved
        self.click(self.button("Open catalogue"))
        self.choose("Save and open")
        self.launch.assert_awaited_once()
        self.key("Arrow Right")
        self.assertEqual(self.description().value, "Second box")

    def test_missing_file_or_application_has_a_clear_error_and_keeps_the_form(self):
        self.open_folder()
        self.launch.side_effect = OSError("No application is associated with this file")
        with self.assertLogs(main.logger, level="WARNING"):
            self.click(self.button("Open catalogue"))
        self.assertIn("default app", self.page.dialogs[-1].content.value)
        self.choose("OK")
        self.launch.reset_mock()
        self.path.unlink()
        self.click(self.button("Open catalogue"))
        self.assertIn("could not be found", self.page.dialogs[-1].content.value)
        self.launch.assert_not_awaited()
        self.assertEqual(self.description().value, "First box")
        self.assertFalse(self.button("Open catalogue").disabled)

    def test_open_failure_keeps_unsaved_edits_and_can_be_retried(self):
        self.open_folder()
        self.description().value = "Unsaved correction"
        self.launch.side_effect = RuntimeError("No association")
        self.click(self.button("Open catalogue"))
        with self.assertLogs(main.logger, level="WARNING"):
            self.choose("Open saved version")
        self.assertEqual(self.description().value, "Unsaved correction")
        self.assertFalse(self.button("Open catalogue").disabled)
        self.choose("OK")
        self.launch.side_effect = None
        self.click(self.button("Open catalogue"))
        self.choose("Open saved version")
        self.assertEqual(self.launch.await_count, 2)

    def test_repeated_open_clicks_do_not_launch_multiple_apps(self):
        self.open_folder()

        async def exercise():
            started = asyncio.Event()
            release = asyncio.Event()

            async def hold(*args, **kwargs):
                started.set()
                await release.wait()

            self.launch.side_effect = hold
            self.click(self.button("Open catalogue"))
            await asyncio.wait_for(started.wait(), timeout=1)
            self.assertTrue(self.button("Open catalogue").disabled)
            self.click(self.button("Open catalogue"))
            self.key("Arrow Right")
            self.assertEqual(self.description().value, "First box")
            release.set()
            await asyncio.gather(*self.page.tasks)
            self.launch.assert_awaited_once()
            self.assertFalse(self.button("Open catalogue").disabled)

        asyncio.run(exercise())

    def test_navigation_still_prompts_and_saves(self):
        self.open_folder()
        self.description().value = "Corrected first box"
        self.click(self.button("Next"))
        self.choose("Save")
        self.assertEqual(self.description().value, "Second box")
        self.assertEqual(main.load_rows(self.path)[1][0]["Description"], "Corrected first box")

    def test_reopening_dirty_folder_can_be_cancelled(self):
        self.open_folder()
        self.description().value = "Unsaved correction"
        self.open_folder()
        self.assertEqual(self.page.dialogs[-1].title.value, "Unsaved changes")
        self.choose("Cancel")
        self.assertEqual(self.description().value, "Unsaved correction")

    def test_folder_switch_save_failure_keeps_current_edits(self):
        self.open_folder()
        self.description().value = "Unsaved correction"
        self.open_folder()
        with patch.object(main, "save_rows", side_effect=main.CatalogueError("Disk full")):
            self.choose("Save")
        self.assertEqual(self.description().value, "Unsaved correction")
        self.assertIn("Disk full", self.page.dialogs[-1].content.value)

    def test_failed_save_stays_dirty_and_discard_restores_original(self):
        self.open_folder()
        self.description().value = "Unsaved correction"
        with patch.object(main, "save_rows", side_effect=main.CatalogueError("Disk full")):
            self.click(self.button("Save"))
        self.choose("OK")
        self.click(self.button("Next"))
        self.assertEqual(self.page.dialogs[-1].title.value, "Unsaved changes")
        self.choose("Cancel")
        self.click(self.button("Discard changes"))
        self.choose("Discard changes")
        self.assertEqual(self.description().value, "First box")

    def test_long_error_details_can_scroll_without_losing_unsaved_edits(self):
        self.open_folder()
        self.description().value = "Unsaved correction"
        message = "Long error detail. " * 250
        with patch.object(main, "save_rows", side_effect=main.CatalogueError(message)):
            self.click(self.button("Save"))
        dialog = self.page.dialogs[-1]
        self.assertTrue(dialog.scrollable)
        self.assertEqual(dialog.content.value, message)
        self.choose("OK")
        self.assertEqual(self.description().value, "Unsaved correction")

    def test_failed_discard_keeps_both_records(self):
        self.open_folder()
        self.click(self.button("Discard box"))
        with patch.object(main, "save_rows", side_effect=main.CatalogueError("Permission denied")):
            self.choose("Discard box")
        self.choose("OK")
        self.assertEqual(self.description().value, "First box")
        self.click(self.button("Next"))
        self.assertEqual(self.description().value, "Second box")
        self.assertEqual(len(main.load_rows(self.path)[1]), 2)

    def test_external_edit_blocks_save(self):
        self.open_folder()
        self.description().value = "GUI correction"
        self.path.write_text(self.path.read_text().replace("Second box", "External correction"))
        self.click(self.button("Save"))
        self.assertIsInstance(self.page.dialogs[-1], ft.AlertDialog)
        self.assertEqual(self.description().value, "GUI correction")
        self.assertEqual(main.load_rows(self.path)[1][1]["Description"], "External correction")

    def test_bad_folder_preserves_existing_form(self):
        self.open_folder()
        other = self.folder / "other"
        other.mkdir()
        (other / "catalogue.xlsx").write_text("Not a workbook")
        self.open_folder(other)
        self.assertEqual(self.description().value, "First box")
        self.assertIsInstance(self.page.dialogs[-1], ft.AlertDialog)

    def test_malformed_workbook_errors_preserve_the_existing_dirty_form(self):
        import openpyxl

        self.open_folder()
        self.description().value = "Unsaved correction"
        other = self.folder / "other"
        other.mkdir()
        (other / "catalogue.xlsx").write_bytes(b"Malformed workbook")
        for error in (TypeError("Invalid sheet ID"), IndexError("Invalid style index")):
            with self.subTest(error=error):
                self.open_folder(other)
                with patch.object(openpyxl, "load_workbook", side_effect=error):
                    self.choose("Don't save")
                self.assertEqual(self.page.dialogs[-1].title.value, "Couldn't open folder")
                self.assertEqual(self.description().value, "Unsaved correction")
                self.choose("OK")

    def test_oversized_workbook_edit_remains_dirty_until_corrected_and_retried(self):
        import openpyxl

        other = self.folder / "excel"
        other.mkdir()
        path = other / "catalogue.xlsx"
        workbook = openpyxl.Workbook()
        workbook.active.append(["Description", main.REVIEWED_FIELD])
        workbook.active.append(["Original description", False])
        workbook.save(path)
        workbook.close()
        self.open_folder(other)
        original = path.read_bytes()
        self.description().value = "x" * 32_768
        self.click(self.button("Save"))
        self.assertIn("32,767", self.page.dialogs[-1].content.value)
        self.assertEqual(path.read_bytes(), original)
        self.assertEqual(len(self.description().value), 32_768)
        self.choose("OK")
        self.page.window.on_event(SimpleNamespace(type=ft.WindowEventType.CLOSE))
        self.assertEqual(self.page.dialogs[-1].title.value, "Unsaved changes")
        self.choose("Cancel")
        self.description().value = "Shortened correction"
        self.click(self.button("Save"))
        self.assertEqual(main.load_rows(path)[1][0]["Description"], "Shortened correction")

    def test_added_flags_get_saved_when_header_was_missing(self):
        self.write_catalogue(include_flags=False)
        self.open_folder()
        menu = self.button("+ Add flag")
        self.click(menu.items[0])
        self.click(self.button("Save"))
        fields, rows = main.load_rows(self.path)
        self.assertIn(main.FLAGS_FIELD, fields)
        self.assertEqual(rows[0][main.FLAGS_FIELD], main.DUPLICATE_FLAG)

    def test_closing_dirty_window_can_be_cancelled(self):
        self.open_folder()
        self.description().value = "Unsaved correction"
        self.page.window.on_event(SimpleNamespace(type=ft.WindowEventType.CLOSE))
        self.choose("Cancel")
        self.page.window.destroy.assert_not_awaited()
        self.assertEqual(self.description().value, "Unsaved correction")

    def test_clean_window_closes(self):
        self.page.window.on_event(SimpleNamespace(type=ft.WindowEventType.CLOSE))
        self.page.window.destroy.assert_awaited_once()

    def test_failed_save_on_close_keeps_window_open_and_edits(self):
        self.open_folder()
        self.description().value = "Unsaved correction"
        self.page.window.on_event(SimpleNamespace(type=ft.WindowEventType.CLOSE))
        with patch.object(main, "save_rows", side_effect=main.CatalogueError("Disk full")):
            self.choose("Save")
        self.page.window.destroy.assert_not_awaited()
        self.assertEqual(self.description().value, "Unsaved correction")
        self.assertEqual(main.load_rows(self.path)[1][0]["Description"], "First box")
        self.assertIn("Disk full", self.page.dialogs[-1].content.value)

    def test_discarding_all_records_leaves_safe_empty_state(self):
        self.open_folder()
        for _ in range(2):
            self.click(self.button("Discard box"))
            self.choose("Discard box")
        self.assertEqual(main.load_rows(self.path)[1], [])
        self.assertEqual(self.description().value, "")
        self.assertTrue(self.description().disabled)
        for label in ("Save", "Discard box", "Prev", "Next"):
            self.assertTrue(self.button(label).disabled)
        with patch.object(main, "save_rows") as save:
            for label in ("Save", "Discard box", "Discard changes"):
                self.click(self.button(label))
            save.assert_not_called()
        self.assertEqual(self.page.dialogs, [])

    def test_save_and_discard_ignore_queued_clicks_without_an_active_record(self):
        for state in ("no catalogue", "all reviewed"):
            with self.subTest(state=state):
                if state == "all reviewed":
                    self.write_catalogue(reviewed=True)
                    self.open_folder()
                with patch.object(main, "save_rows") as save:
                    for label in ("Save", "Discard box", "Discard changes"):
                        self.click(self.button(label))
                    save.assert_not_called()
                self.assertEqual(self.page.dialogs, [])

    def test_reviewed_catalogue_has_clear_completion_and_revisit_action(self):
        self.write_catalogue(reviewed=True)
        self.open_folder()
        messages = [c.value for c in self.controls() if isinstance(c, ft.Text)]
        self.assertTrue(any("All 2 records reviewed" in (message or "") for message in messages))
        self.click(self.button("View records"))
        self.assertEqual(self.description().value, "First box")

    def test_unreadable_crop_does_not_stop_review(self):
        self.path.write_text(self.path.read_text().replace("First box,,", "First box,crop.png,"))
        (self.folder / "crop.png").write_bytes(b"image")
        with patch.object(Path, "read_bytes", side_effect=PermissionError("Unreadable image")):
            self.open_folder()
        self.assertEqual(self.description().value, "First box")
        messages = [c.value for c in self.controls() if isinstance(c, ft.Text)]
        self.assertTrue(any("image" in (message or "").lower() for message in messages))

    def test_invalid_crop_path_does_not_stop_review(self):
        self.path.write_text(self.path.read_text().replace("First box,,", "First box,invalid\x00path,"))
        self.open_folder()
        self.assertEqual(self.description().value, "First box")
        messages = [c.value for c in self.controls() if isinstance(c, ft.Text)]
        self.assertIn("Couldn't read this image", messages)

    @unittest.skipUnless(hasattr(os, "mkfifo"), "Named pipes require POSIX")
    def test_fifo_crop_is_rejected_without_opening_or_blocking(self):
        self.path.write_text(self.path.read_text().replace("First box,,", "First box,crop.png,"))
        os.mkfifo(self.folder / "crop.png")
        with patch.object(Path, "read_bytes", side_effect=AssertionError("Do not read a FIFO")) as read:
            self.open_folder()
            read.assert_not_called()
        self.assertEqual(self.description().value, "First box")
        messages = [c.value for c in self.controls() if isinstance(c, ft.Text)]
        self.assertIn("Couldn't read this image: the path is not a regular file", messages)

    def test_window_is_centred_shown_and_remains_usable_if_centring_fails(self):
        self.page.window.center.assert_awaited_once()
        self.assertTrue(self.page.window.visible)
        self.assertEqual((self.page.window.width, self.page.window.height), (1000, 650))
        self.page = TestPage()
        self.page.window.center.side_effect = RuntimeError("Window manager unavailable")
        with self.assertLogs(main.logger, level="WARNING"):
            asyncio.run(main.page_main(self.page))
        self.assertTrue(self.page.window.visible)
        self.open_folder()
        self.assertEqual(self.description().value, "First box")

    def test_theme_toggle_remembers_choice_without_losing_edits(self):
        self.open_folder()
        self.description().value = "Unsaved correction"
        toggle = next(c for c in self.controls() if isinstance(c, ft.IconButton))
        self.click(toggle)
        self.assertEqual(self.page.theme_mode, ft.ThemeMode.DARK)
        self.assertEqual(toggle.tooltip, "Switch to light mode")
        self.assertEqual(self.preferences[main.THEME_PREFERENCE], "dark")
        self.assertEqual(self.description().value, "Unsaved correction")
        self.click(self.button("Next"))
        self.assertEqual(self.page.dialogs[-1].title.value, "Unsaved changes")
        self.choose("Cancel")
        self.page = TestPage()
        asyncio.run(main.page_main(self.page))
        self.assertEqual(self.page.theme_mode, ft.ThemeMode.DARK)

    def test_theme_uses_system_brightness_until_an_explicit_choice(self):
        self.page.platform_brightness = ft.Brightness.DARK
        self.page.on_platform_brightness_change(None)
        toggle = next(c for c in self.controls() if isinstance(c, ft.IconButton))
        self.assertEqual(toggle.tooltip, "Switch to light mode")
        self.click(toggle)
        self.assertEqual(self.page.theme_mode, ft.ThemeMode.LIGHT)
        self.page.on_platform_brightness_change(None)
        self.assertEqual(self.page.theme_mode, ft.ThemeMode.LIGHT)

    def test_preference_failures_do_not_block_the_app_or_toggle(self):
        self.get_preference.side_effect = RuntimeError("Unavailable storage")
        self.page = TestPage()
        with self.assertLogs(main.logger, level="WARNING"):
            asyncio.run(main.page_main(self.page))
        self.assertTrue(self.page.window.visible)
        self.assertEqual(self.page.theme_mode, ft.ThemeMode.SYSTEM)
        self.set_preference.side_effect = RuntimeError("Unavailable storage")
        toggle = next(c for c in self.controls() if isinstance(c, ft.IconButton))
        with self.assertLogs(main.logger, level="WARNING"):
            self.click(toggle)
        self.assertEqual(self.page.theme_mode, ft.ThemeMode.DARK)
        self.assertFalse(toggle.disabled)
        self.assertIn("this session", self.page.dialogs[-1].content.value)

    def test_theme_failure_cannot_keep_discard_confirmation_open_or_repeat_deletion(self):
        self.open_folder()

        async def exercise():
            started = asyncio.Event()
            release = asyncio.Event()

            async def fail_later(key, value):
                started.set()
                await release.wait()
                raise RuntimeError("Unavailable storage")

            self.set_preference.side_effect = fail_later
            toggle = next(c for c in self.controls() if isinstance(c, ft.IconButton))
            task = asyncio.create_task(toggle.on_click(None))
            await asyncio.wait_for(started.wait(), timeout=1)
            self.click(self.button("Discard box"))
            confirmation = self.page.dialogs[-1]
            confirm_button = next(c for c in confirmation.actions if c.content == "Discard box")

            with self.assertLogs(main.logger, level="WARNING"):
                release.set()
                await task
            notification = self.page.dialogs[-1]
            self.assertIsInstance(notification, ft.SnackBar)
            self.click(confirm_button)
            self.assertFalse(confirmation.open)
            self.assertTrue(notification.open)
            self.assertEqual(len(main.load_rows(self.path)[1]), 1)

            # A second click can already be queued while the dialog is closing.
            self.click(confirm_button)
            self.assertEqual(len(main.load_rows(self.path)[1]), 1)

        asyncio.run(exercise())

    def test_resizing_retains_unsaved_form_values_and_flags(self):
        self.open_folder()
        description = self.description()
        description.value = "Unsaved correction"
        self.click(self.button("+ Add flag").items[0])
        for width in (640, 1000, 800):
            self.page.on_resize(SimpleNamespace(width=width))
            self.assertIs(self.description(), description)
            self.assertEqual(description.value, "Unsaved correction")
            chips = [c for c in self.controls() if isinstance(c, ft.Chip)]
            self.assertEqual([c.data for c in chips], [main.DUPLICATE_FLAG])
        self.click(self.button("Save"))
        row = main.load_rows(self.path)[1][0]
        self.assertEqual(row["Description"], "Unsaved correction")
        self.assertEqual(row[main.FLAGS_FIELD], main.DUPLICATE_FLAG)

    def test_short_workspaces_scroll_all_content_without_changing_live_parents(self):
        self.open_folder()
        description = self.description()
        description.value = "Unsaved correction"
        self.click(self.button("+ Add flag").items[0])
        root = self.page.controls[0]
        header, content, footer = root.controls
        body = content.controls[0]
        form, image = body.controls
        for width, height in ((640, 288), (1000, 288), (640, 240), (1000, 618), (640, 448)):
            with self.subTest(width=width, height=height):
                self.page.on_resize(SimpleNamespace(width=width, height=height))
                self.assertIs(root.controls[0], header)
                self.assertIs(root.controls[1], content)
                self.assertIs(root.controls[2], footer)
                self.assertIs(content.controls[0], body)
                self.assertIs(body.controls[0], form)
                self.assertIs(body.controls[1], image)
                self.assertIs(self.description(), description)
                self.assertEqual(description.value, "Unsaved correction")
                self.assertEqual(root.scroll, ft.ScrollMode.AUTO if height < 400 else None)
                if height < 400:
                    self.assertIsNone(content.expand)
                    self.assertTrue(content.tight)
                    self.assertIsNone(content.scroll)
                    self.assertIsNone(form.scroll)
                    self.assertIsNone(body.expand)
                else:
                    self.assertTrue(content.expand)
                    self.assertEqual(content.scroll, ft.ScrollMode.AUTO if width < 850 else None)
        self.click(self.button("Save"))
        row = main.load_rows(self.path)[1][0]
        self.assertEqual(row["Description"], "Unsaved correction")
        self.assertEqual(row[main.FLAGS_FIELD], main.DUPLICATE_FLAG)

    def test_long_field_names_use_one_line_and_expose_the_full_name(self):
        name = "Long catalogue column " * 15
        with self.path.open("w", newline="", encoding="utf-8") as output:
            writer = csv.writer(output)
            writer.writerow([name, "Description"])
            writer.writerow(["Original value", "First box"])
        self.open_folder()
        field = next(c for c in self.controls() if isinstance(c, ft.TextField) and getattr(c.label, "value", c.label) == name)
        self.assertIsInstance(field.label, ft.Text)
        self.assertTrue(field.label.no_wrap)
        self.assertEqual(field.label.max_lines, 1)
        self.assertEqual(field.label.overflow, ft.TextOverflow.ELLIPSIS)
        self.assertEqual(field.tooltip, name)
        field.value = "Corrected value"
        self.click(self.button("Save"))
        self.assertEqual(main.load_rows(self.path)[1][0][name], "Corrected value")

    def test_failed_save_can_be_retried_without_losing_edits(self):
        self.open_folder()
        self.description().value = "Correction"
        with patch("main.os.replace", side_effect=PermissionError("File locked")):
            self.click(self.button("Save"))
        self.assertEqual(self.description().value, "Correction")
        self.choose("OK")
        self.click(self.button("Save"))
        self.assertEqual(main.load_rows(self.path)[1][0]["Description"], "Correction")

    def test_uppercase_reviewed_and_selected_filename_are_shown_correctly(self):
        self.write_catalogue(reviewed="TRUE")
        renamed = self.path.with_name("Catalogue.CSV")
        self.path.rename(renamed)
        self.open_folder()
        messages = [c.value for c in self.controls() if isinstance(c, ft.Text)]
        self.assertIn("All 2 records reviewed", messages)
        self.assertIn("Editing Catalogue.CSV", messages)

    def test_batch_with_hyperlink_formulas_shows_seven_images_and_two_empty_records(self):
        image = base64.b64decode("iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mNk+A8AAQUBAScY42YAAAAASUVORK5CYII=")
        crops = self.folder / "Box Images"
        crops.mkdir()
        fields = ["Description", main.IMAGE_LINK_FIELD, main.REVIEWED_FIELD]
        with self.path.open("w", newline="", encoding="utf-8") as output:
            writer = csv.writer(output)
            writer.writerow(fields)
            for index in range(9):
                link = ""
                if index not in (1, 2):
                    (crops / f"Box {index}.png").write_bytes(image)
                    link = f'=HYPERLINK("Box Images/Box {index}.png","Box {index}")'
                writer.writerow([f"Record {index}", link, "False"])
        original_links = [row[main.IMAGE_LINK_FIELD] for row in main.load_rows(self.path)[1]]
        self.open_folder()
        image_count = empty_count = 0
        for index in range(9):
            previews = [c for c in self.controls() if isinstance(c, ft.Image)]
            if previews:
                image_count += 1
                self.assertEqual(previews[0].src, image)
            else:
                empty_count += 1
                self.assertIn("No image", [c.value for c in self.controls() if isinstance(c, ft.Text)])
            if index < 8:
                self.click(self.button("Next"))
        self.assertEqual((image_count, empty_count), (7, 2))
        self.description().value = "Correction"
        self.click(self.button("Save"))
        self.assertEqual([row[main.IMAGE_LINK_FIELD] for row in main.load_rows(self.path)[1]], original_links)

    def test_empty_and_damaged_images_have_distinct_fallbacks(self):
        self.path.write_text(self.path.read_text().replace("First box,,", "First box,crop.png,"))
        image = self.folder / "crop.png"
        image.touch()
        self.open_folder()
        self.assertTrue(any("file is empty" in (c.value or "") for c in self.controls() if isinstance(c, ft.Text)))
        image.write_bytes(b"invalid image")
        self.open_folder()
        preview = next(c for c in self.controls() if isinstance(c, ft.Image))
        self.assertIn("damaged or unsupported", preview.error_content.value)


if __name__ == "__main__":
    unittest.main()
